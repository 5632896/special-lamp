from __future__ import annotations

import csv
import copy
import hashlib
import io
import json
import os
import re
import ssl
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Literal

import simplemma
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel
from pypdf import PdfReader
try:
    from .runtime_paths import DATA_DIR, FRONTEND_DIST_DIR
    from .storage import (
        append_chat_message, create_chat_session, delete_chat_session, get_chat_session,
        clear_review_schedules, get_due_review_schedules, get_review_schedule, initialize as initialize_storage,
        list_chat_sessions, load_ai_cache as sqlite_load_ai_cache,
        load_libraries as sqlite_load_libraries, load_setting as sqlite_load_setting,
        load_user_status as sqlite_load_user_status, migrate_json_once,
        rename_chat_session, save_review_schedule,
        save_ai_cache as sqlite_save_ai_cache, save_libraries as sqlite_save_libraries,
        save_setting as sqlite_save_setting, save_user_status as sqlite_save_user_status,
    )
except ImportError:
    from runtime_paths import DATA_DIR, FRONTEND_DIST_DIR
    from storage import (
        append_chat_message, create_chat_session, delete_chat_session, get_chat_session,
        clear_review_schedules, get_due_review_schedules, get_review_schedule, initialize as initialize_storage,
        list_chat_sessions, load_ai_cache as sqlite_load_ai_cache,
        load_libraries as sqlite_load_libraries, load_setting as sqlite_load_setting,
        load_user_status as sqlite_load_user_status, migrate_json_once,
        rename_chat_session, save_review_schedule,
        save_ai_cache as sqlite_save_ai_cache, save_libraries as sqlite_save_libraries,
        save_setting as sqlite_save_setting, save_user_status as sqlite_save_user_status,
    )

LIBRARIES_PATH = DATA_DIR / "vocab_libraries.json"
USER_STATUS_PATH = DATA_DIR / "user_vocab_status.json"
AI_SETTINGS_PATH = DATA_DIR / "ai_settings.json"
AI_CACHE_PATH = DATA_DIR / "ai_cache.json"

OLD_BOOK_VOCAB_PATH = DATA_DIR / "book_vocab.json"
OLD_EXTRA_VOCAB_PATH = DATA_DIR / "extra_vocab.json"

EXTRA_LIBRARY_ID = "extra"
BASIC_LIBRARY_ID = "basic_whitelist"
MAX_ANALYZE_WORDS = 5000
MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
SUPPORTED_UPLOAD_EXTENSIONS = {".csv", ".txt", ".xlsx", ".docx", ".pdf"}

MarkState = Literal["known", "fuzzy", "unknown"]
AnalysisMode = Literal["friendly", "strict"]

WORD_COLUMN_ALIASES = {
    "lemma",
    "word",
    "words",
    "term",
    "vocab",
    "vocabulary",
    "english",
    "english_word",
    "english_words",
    "单词",
    "词",
    "词汇",
    "词元",
    "原形",
    "英文",
    "英语",
}
UNIT_COLUMN_ALIASES = {
    "unit",
    "单元",
}
LESSON_COLUMN_ALIASES = {"lesson", "course", "class", "chapter", "课程", "课", "章节"}
PAGE_COLUMN_ALIASES = {
    "page",
    "page_no",
    "page_number",
    "页码",
    "页",
}
FREQUENCY_COLUMN_ALIASES = {
    "frequency",
    "freq",
    "高频低频",
    "频率",
}
IN_SYLLABUS_COLUMN_ALIASES = {
    "in_syllabus",
    "insyllabus",
    "书内",
    "教材内",
    "是否书内",
    "是否教材内",
}
OVERSYLLABUS_COLUMN_ALIASES = {
    "oversyllabus",
    "out_of_syllabus",
    "超纲",
    "是否超纲",
}
MEANING_COLUMN_ALIASES = {
    "meaning",
    "translation",
    "chinese",
    "cn",
    "中文",
    "释义",
    "意思",
    "含义",
}
POS_COLUMN_ALIASES = {"pos", "part_of_speech", "word_class", "词性"}
IPA_COLUMN_ALIASES = {"ipa", "phonetic", "phonetics", "pronunciation", "音标", "发音"}
AUDIO_URL_COLUMN_ALIASES = {"audio_url", "audio", "audio_url", "sound_url", "音频", "音频链接", "发音链接"}
SERIAL_COLUMN_ALIASES = {"serial", "serial_no", "serial_number", "no", "number", "index", "序号", "单词序号", "编号"}

IMPORT_MAPPING_FIELDS = (
    ("lemma", "单词（必填）"),
    ("meaning", "释义"),
    ("unit", "单元"),
    ("lesson", "课程/章节"),
    ("page", "页码"),
    ("frequency", "频率"),
    ("pos", "词性"),
    ("ipa", "音标"),
    ("audio_url", "音频链接"),
    ("serial", "序号"),
    ("in_syllabus", "是否书内"),
)

BUILTIN_STOPWORDS = {
    "the", "a", "an", "and", "or", "but", "so", "because",
    "i", "you", "he", "she", "it", "we", "they", "me", "him", "her", "us", "them",
    "my", "your", "his", "its", "our", "their",
    "is", "am", "are", "was", "were", "be", "been", "being",
    "do", "does", "did", "have", "has", "had",
    "can", "could", "will", "would", "shall", "should", "may", "might", "must",
    "to", "of", "in", "on", "at", "for", "with", "from", "by", "about", "as", "into", "over", "after", "before",
    "not", "no", "yes"
}

CONTRACTION_DIRECT_MAP = {
    "don't": "do",
    "doesn't": "do",
    "didn't": "do",
    "can't": "can",
    "won't": "will",
    "i'm": "i",
    "you're": "you",
    "we're": "we",
    "they're": "they",
    "he's": "he",
    "she's": "she",
    "it's": "it",
}

AI_PROVIDER_PRESETS = [
    {
        "id": "openai-compatible",
        "name": "OpenAI Compatible",
        "baseUrlExample": "https://api.openai.com/v1",
        "modelExample": "gpt-4o-mini",
        "note": "通用 OpenAI 兼容接口",
    },
    {
        "id": "openai",
        "name": "OpenAI",
        "baseUrlExample": "https://api.openai.com/v1",
        "modelExample": "gpt-4o-mini",
        "note": "官方 OpenAI 接口",
    },
    {
        "id": "deepseek",
        "name": "DeepSeek",
        "baseUrlExample": "https://api.deepseek.com/v1",
        "modelExample": "deepseek-chat",
        "note": "DeepSeek 官方兼容接口",
    },
    {
        "id": "moonshot",
        "name": "Moonshot / Kimi",
        "baseUrlExample": "https://api.moonshot.cn/v1",
        "modelExample": "moonshot-v1-8k",
        "note": "Kimi / Moonshot 兼容接口",
    },
    {
        "id": "dashscope",
        "name": "DashScope / Qwen",
        "baseUrlExample": "https://dashscope.aliyuncs.com/compatible-mode/v1",
        "modelExample": "qwen-plus",
        "note": "阿里云百炼 / 通义千问兼容接口",
    },
    {
        "id": "siliconflow",
        "name": "SiliconFlow",
        "baseUrlExample": "https://api.siliconflow.cn/v1",
        "modelExample": "Qwen/Qwen2.5-7B-Instruct",
        "note": "SiliconFlow 兼容接口",
    },
    {
        "id": "groq",
        "name": "Groq",
        "baseUrlExample": "https://api.groq.com/openai/v1",
        "modelExample": "llama-3.1-8b-instant",
        "note": "Groq OpenAI 兼容接口",
    },
    {
        "id": "together",
        "name": "Together AI",
        "baseUrlExample": "https://api.together.xyz/v1",
        "modelExample": "meta-llama/Llama-3.1-8B-Instruct-Turbo",
        "note": "Together AI 兼容接口",
    },
    {
        "id": "openrouter",
        "name": "OpenRouter",
        "baseUrlExample": "https://openrouter.ai/api/v1",
        "modelExample": "openai/gpt-4o-mini",
        "note": "OpenRouter 多模型路由",
    },
    {
        "id": "ollama",
        "name": "Ollama",
        "baseUrlExample": "http://127.0.0.1:11434/v1",
        "modelExample": "qwen2.5:7b",
        "note": "本地 Ollama OpenAI 兼容接口",
    },
]
DISPLAY_FIELD_GROUPS = {
    "tokenFields": [
        {"key": "meaning", "label": "释义", "source": "meaning", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "unit", "label": "单元", "source": "unit", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "page", "label": "页码", "source": "page", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "frequency", "label": "频率", "source": "frequencyText", "enabled": True, "displayType": "text", "showEmpty": False},
    ],
    "tooltipFields": [
        {"key": "lemma", "label": "词形", "source": "lemma", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "basicText", "label": "基础词", "source": "basicText", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "outsideText", "label": "词库外", "source": "outsideText", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "sourceText", "label": "来源", "source": "sourceText", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "unit", "label": "单元", "source": "unit", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "page", "label": "页码", "source": "page", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "frequencyText", "label": "频率", "source": "frequencyText", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "level", "label": "掌握等级", "source": "levelLabel", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "meaning", "label": "词义", "source": "meaning", "enabled": True, "displayType": "text", "showEmpty": False},
    ],
    "focusFields": [
        {"key": "meaning", "label": "释义", "source": "meaning", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "libraryName", "label": "词库", "source": "libraryName", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "unit", "label": "单元", "source": "unit", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "page", "label": "页码", "source": "page", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "frequency", "label": "频率", "source": "frequencyText", "enabled": True, "displayType": "text", "showEmpty": False},
    ],
    "detailFields": [
        {"key": "meaning", "label": "释义", "source": "meaning", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "unit", "label": "单元", "source": "unit", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "page", "label": "页码", "source": "page", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "frequency", "label": "频率", "source": "frequencyText", "enabled": True, "displayType": "text", "showEmpty": False},
    ],
    "wordListFields": [
        {"key": "meaning", "label": "释义", "source": "meaning", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "ipa", "label": "音标", "source": "ipa", "enabled": True, "displayType": "text", "showEmpty": False},
        {"key": "unit", "label": "单元", "source": "unit", "enabled": True, "displayType": "text", "showEmpty": False},
    ],
}


def default_display_config(library: dict) -> dict:
    config = json.loads(json.dumps(DISPLAY_FIELD_GROUPS, ensure_ascii=False))
    config["focusFieldLimit"] = 5
    config.update({"listMode": "recall", "listLayout": "list", "rushMode": "recall",
                   "detailPlacement": "below", "allowMultipleExpanded": True})
    custom_names = []
    for entry in library.get("entries", []):
        fields = entry.get("customFields") if isinstance(entry.get("customFields"), dict) else {}
        for name in fields:
            if name not in custom_names:
                custom_names.append(name)
    custom_fields = [
        {"key": f"custom:{name}", "label": name, "source": f"customFields.{name}", "enabled": True, "displayType": "text", "showEmpty": False}
        for name in custom_names
    ]
    config["detailFields"].extend(custom_fields)
    for definition in library.get("fieldDefinitions", []):
        if not isinstance(definition, dict):
            continue
        key = safe_text(definition.get("key"))
        label = safe_text(definition.get("label")) or key
        if key and not any(field.get("source") == f"customFields.{key}" for field in config["detailFields"]):
            config["detailFields"].append({"key": f"custom:{key}", "label": label, "source": f"customFields.{key}", "enabled": True, "displayType": "text", "showEmpty": False})
    return config


def normalized_display_config(library: dict) -> dict:
    saved = library.get("displayConfig")
    if not isinstance(saved, dict) or not saved:
        return default_display_config(library)
    defaults = default_display_config(library)
    result = {}
    for group in DISPLAY_FIELD_GROUPS:
        fields = saved.get(group)
        result[group] = fields if isinstance(fields, list) and (fields or group == "wordListFields") else defaults[group]
    try:
        field_limit = int(saved.get("focusFieldLimit", defaults["focusFieldLimit"]))
    except (TypeError, ValueError):
        field_limit = defaults["focusFieldLimit"]
    result["focusFieldLimit"] = max(1, min(12, field_limit))
    result["listMode"] = saved.get("listMode") if saved.get("listMode") in {"browse", "recall"} else defaults["listMode"]
    result["listLayout"] = saved.get("listLayout") if saved.get("listLayout") in {"list", "tiles"} else defaults["listLayout"]
    result["rushMode"] = saved.get("rushMode") if saved.get("rushMode") in {"browse", "recall"} else defaults["rushMode"]
    result["detailPlacement"] = saved.get("detailPlacement") if saved.get("detailPlacement") in {"below", "floating"} else defaults["detailPlacement"]
    result["allowMultipleExpanded"] = bool(saved.get("allowMultipleExpanded", defaults["allowMultipleExpanded"]))
    return result


def validate_display_config(library: dict, config: dict) -> dict:
    allowed = {key for key, _ in IMPORT_MAPPING_FIELDS} | {
        "libraryName", "frequencyText", "sourceText", "outsideText", "level", "levelLabel", "meaning", "unit", "page", "lemma", "pos", "ipa", "audioUrl",
        "basicText", "count", "fuzzyCount", "unknownCount", "suggestedLevel"
    }
    custom_names = {
        name
        for entry in library.get("entries", [])
        for name in (entry.get("customFields", {}) if isinstance(entry.get("customFields"), dict) else {})
    }
    custom_names.update(
        safe_text(definition.get("key"))
        for definition in library.get("fieldDefinitions", [])
        if isinstance(definition, dict) and safe_text(definition.get("key"))
    )
    normalized = {}
    for group in DISPLAY_FIELD_GROUPS:
        # Older clients have no wordListFields; keep their saves compatible.
        raw_fields = config.get(group, default_display_config(library)[group] if group == "wordListFields" else [])
        if not isinstance(raw_fields, list):
            raise HTTPException(status_code=400, detail=f"{group} 必须是字段数组。")
        seen = set()
        fields = []
        for index, field in enumerate(raw_fields):
            if not isinstance(field, dict):
                continue
            source = safe_text(field.get("source"))
            label = safe_text(field.get("label")) or source
            key = safe_text(field.get("key")) or f"{group}:{index}"
            custom_name = source.removeprefix("customFields.") if source.startswith("customFields.") else ""
            if source not in allowed and custom_name not in custom_names:
                raise HTTPException(status_code=400, detail=f"字段来源无效：{source}")
            if key in seen:
                raise HTTPException(status_code=400, detail=f"{group} 存在重复字段。")
            seen.add(key)
            fields.append({"key": key, "label": label[:80], "source": source, "enabled": bool(field.get("enabled", True)), "displayType": "text", "showEmpty": bool(field.get("showEmpty", False))})
        normalized[group] = fields
    try:
        focus_field_limit = int(config.get("focusFieldLimit", 5))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="重点词清单字段上限必须是数字。")
    normalized["focusFieldLimit"] = max(1, min(12, focus_field_limit))
    normalized["listMode"] = config.get("listMode") if config.get("listMode") in {"browse", "recall"} else "recall"
    list_layout = config.get("listLayout", "list")
    if list_layout not in {"list", "tiles"}:
        raise HTTPException(status_code=400, detail="单词列表布局必须是 list 或 tiles。")
    normalized["listLayout"] = list_layout
    normalized["rushMode"] = config.get("rushMode") if config.get("rushMode") in {"browse", "recall"} else "recall"
    normalized["detailPlacement"] = config.get("detailPlacement") if config.get("detailPlacement") in {"below", "floating"} else "below"
    normalized["allowMultipleExpanded"] = bool(config.get("allowMultipleExpanded", True))
    return normalized


def available_display_fields(library: dict) -> list[dict]:
    # Keep analysis display choices independent from import-column mapping. The
    # `present` flag is only a hint for the UI's automatic current-library match.
    entries = library.get("entries", []) if isinstance(library.get("entries"), list) else []
    present = set()
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        if safe_text(entry.get("meaning")):
            present.add("meaning")
        if safe_text(entry.get("unit")):
            present.add("unit")
        if entry.get("page") not in (None, ""):
            present.add("page")
        if safe_text(entry.get("frequency")) or safe_text(entry.get("frequencyText")):
            present.add("frequencyText")
        for source, keys in {
            "pos": ("pos", "partOfSpeech"),
            "ipa": ("ipa", "phonetic"),
            "audioUrl": ("audioUrl", "audio_url"),
        }.items():
            if any(safe_text(entry.get(key)) for key in keys):
                present.add(source)
        custom_fields = entry.get("customFields") if isinstance(entry.get("customFields"), dict) else {}
        for name, value in custom_fields.items():
            if safe_text(value):
                present.add(f"customFields.{name}")

    fields = [
        {"source": "meaning", "label": "释义", "present": "meaning" in present},
        {"source": "unit", "label": "单元", "present": "unit" in present},
        {"source": "page", "label": "页码", "present": "page" in present},
        {"source": "frequencyText", "label": "频率", "present": "frequencyText" in present},
        {"source": "libraryName", "label": "词库", "present": False},
        {"source": "lemma", "label": "词形", "present": False},
        {"source": "sourceText", "label": "来源", "present": False},
        {"source": "outsideText", "label": "词库外说明", "present": False},
        {"source": "level", "label": "掌握等级", "present": False},
        {"source": "levelLabel", "label": "掌握等级名称", "present": False},
        {"source": "pos", "label": "词性", "present": "pos" in present},
        {"source": "ipa", "label": "音标", "present": "ipa" in present},
        {"source": "audioUrl", "label": "音频链接", "present": "audioUrl" in present},
        {"source": "basicText", "label": "基础词状态", "present": False},
        {"source": "count", "label": "出现次数", "present": False},
        {"source": "fuzzyCount", "label": "模糊次数", "present": False},
        {"source": "unknownCount", "label": "未掌握次数", "present": False},
        {"source": "suggestedLevel", "label": "建议等级", "present": False},
    ]
    for definition in library.get("fieldDefinitions", []):
        if isinstance(definition, dict):
            key = safe_text(definition.get("key"))
            source = f"customFields.{key}" if key else ""
            label = safe_text(definition.get("label"))
            if source and label and not any(item["source"] == source for item in fields):
                fields.append({"source": source, "label": label, "present": source in present})
    custom_names = []
    for entry in library.get("entries", []):
        custom_fields = entry.get("customFields") if isinstance(entry.get("customFields"), dict) else {}
        for name in custom_fields:
            if name not in custom_names:
                custom_names.append(name)
    fields.extend({"source": f"customFields.{name}", "label": name, "present": f"customFields.{name}" in present} for name in custom_names)
    return fields

AI_CHAT_TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".html", ".css", ".xml", ".yaml", ".yml", ".log"
}
AI_CHAT_DOC_EXTENSIONS = {".docx", ".pdf", ".xlsx"}
AI_CHAT_SUPPORTED_EXTENSIONS = AI_CHAT_TEXT_EXTENSIONS | AI_CHAT_DOC_EXTENSIONS

POS_ZH_MAP = {
    "noun": "名词",
    "verb": "动词",
    "adjective": "形容词",
    "adverb": "副词",
    "pronoun": "代词",
    "preposition": "介词",
    "conjunction": "连词",
    "interjection": "感叹词",
    "article": "冠词",
    "determiner": "限定词",
    "auxiliary": "助动词",
    "modal": "情态动词",
    "phrasal verb": "短语动词",
}

app = FastAPI(title="English Reading Vocab Diagnosis MVP")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class TokenPayload(BaseModel):
    id: int
    text: str
    isWord: bool


class AnalyzeRequest(BaseModel):
    tokens: list[TokenPayload]
    tokenMarks: dict[str, MarkState] = {}
    selectedLibraryIds: list[str] = []
    includeExtra: bool = True
    analysisMode: AnalysisMode = "friendly"
    ignoreBasicWords: bool = True


class AddExtraWordRequest(BaseModel):
    word: str
    level: int = 1
    frequency: Literal["high", "low"] = "low"
    meaning: str | None = None
    libraryId: str | None = None


class AddBasicWordsRequest(BaseModel):
    words: list[str] = []


class AISettingsPayload(BaseModel):
    provider: str = "openai-compatible"
    baseUrl: str = ""
    model: str = ""
    apiKey: str = ""
    temperature: float = 0.2
    timeoutSeconds: int = 30
    verifySSL: bool = True


class AIWordQueryRequest(BaseModel):
    word: str
    context: str | None = None
    libraryContext: dict | None = None
    forceRefresh: bool = False
    settings: AISettingsPayload | None = None


class AIArticleTranslateRequest(BaseModel):
    text: str
    forceRefresh: bool = False
    settings: AISettingsPayload | None = None


class ImportConfirmRequest(BaseModel):
    libraryName: str = ""
    headers: list[str] = []
    rows: list[list] = []
    mapping: dict[str, str] = {}
    fieldDefinitions: list[dict] = []
    targetType: Literal["main", "basic"] = "main"


class DisplayConfigRequest(BaseModel):
    displayConfig: dict = {}


class AIImportDraftRequest(BaseModel):
    text: str
    settings: AISettingsPayload | None = None


class ReviewMarkRequest(BaseModel):
    lemma: str
    mark: MarkState


class BasicWordsDeleteRequest(BaseModel):
    words: list[str] = []


class LibraryCreateRequest(BaseModel):
    name: str


class LibraryRenameRequest(BaseModel):
    name: str


class LibraryMergeRequest(BaseModel):
    sourceLibraryIds: list[str]
    name: str
    deleteSources: bool = False


class DefaultExtraLibraryRequest(BaseModel):
    libraryId: str


class ChatSessionCreateRequest(BaseModel):
    title: str = "新对话"


class ChatSessionRenameRequest(BaseModel):
    title: str


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def default_ai_settings() -> dict:
    return {
        "provider": "openai-compatible",
        "baseUrl": "",
        "model": "",
        "temperature": 0.2,
        "timeoutSeconds": 30,
        "verifySSL": True,
    }


def ensure_data_files() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if not LIBRARIES_PATH.exists():
        LIBRARIES_PATH.write_text("[]", encoding="utf-8")

    if not USER_STATUS_PATH.exists():
        USER_STATUS_PATH.write_text("{}", encoding="utf-8")

    if not AI_SETTINGS_PATH.exists():
        AI_SETTINGS_PATH.write_text(
            json.dumps(default_ai_settings(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    if not AI_CACHE_PATH.exists():
        AI_CACHE_PATH.write_text(
            json.dumps({"word": {}, "article": {}, "chat": {}}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    migrate_old_vocab_files()
    initialize_storage()
    migrate_json_once(default_ai_settings())


def ensure_data_files_without_migration() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if not LIBRARIES_PATH.exists():
        LIBRARIES_PATH.write_text("[]", encoding="utf-8")

    if not USER_STATUS_PATH.exists():
        USER_STATUS_PATH.write_text("{}", encoding="utf-8")

    if not AI_SETTINGS_PATH.exists():
        AI_SETTINGS_PATH.write_text(
            json.dumps(default_ai_settings(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    if not AI_CACHE_PATH.exists():
        AI_CACHE_PATH.write_text(
            json.dumps({"word": {}, "article": {}, "chat": {}}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


def read_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def write_json(path: Path, data) -> None:
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def load_libraries() -> list[dict]:
    ensure_data_files_without_migration()
    return sqlite_load_libraries()


def save_libraries(libraries: list[dict]) -> None:
    sqlite_save_libraries(libraries)


def load_user_status() -> dict[str, dict]:
    ensure_data_files_without_migration()
    return sqlite_load_user_status()


def save_user_status(status: dict[str, dict]) -> None:
    sqlite_save_user_status(status)


def load_ai_settings() -> dict:
    ensure_data_files_without_migration()
    data = sqlite_load_setting("ai_settings", default_ai_settings())
    if not isinstance(data, dict):
        return default_ai_settings()
    merged = default_ai_settings()
    merged.update({key: value for key, value in data.items() if key != "apiKey"})
    return merged


def save_ai_settings(settings: dict) -> None:
    merged = default_ai_settings()
    merged.update({key: value for key, value in settings.items() if key != "apiKey"})
    sqlite_save_setting("ai_settings", merged)


def environment_api_key() -> str:
    return safe_text(os.getenv("AI_API_KEY")) or safe_text(os.getenv("OPENAI_API_KEY")) or ""


def load_ai_cache() -> dict:
    ensure_data_files_without_migration()
    data = sqlite_load_ai_cache()
    if not isinstance(data, dict):
        return {"word": {}, "article": {}, "chat": {}}
    if not isinstance(data.get("word"), dict):
        data["word"] = {}
    if not isinstance(data.get("article"), dict):
        data["article"] = {}
    if not isinstance(data.get("chat"), dict):
        data["chat"] = {}
    return data


def save_ai_cache(cache: dict) -> None:
    if "word" not in cache or not isinstance(cache["word"], dict):
        cache["word"] = {}
    if "article" not in cache or not isinstance(cache["article"], dict):
        cache["article"] = {}
    if "chat" not in cache or not isinstance(cache["chat"], dict):
        cache["chat"] = {}
    sqlite_save_ai_cache(cache)


def normalize_header(value) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[\s\-/]+", "_", text)
    return text


def safe_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def safe_int(value) -> int | None:
    if value is None:
        return None
    digits = re.findall(r"\d+", str(value))
    if not digits:
        return None
    return int(digits[0])


def normalize_frequency(value) -> str:
    text = (safe_text(value) or "").lower()
    if "high" in text or "高" in text:
        return "high"
    return "low"


def frequency_text(value: str) -> str:
    return "书高频" if value == "high" else "书低频"


def parse_bool(value, default: bool) -> bool:
    if value is None:
        return default
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "是", "书内", "教材内"}:
        return True
    if text in {"0", "false", "no", "n", "否", "超纲"}:
        return False
    return default


def coerce_bool(value, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "on"}:
        return True
    if text in {"0", "false", "no", "n", "off"}:
        return False
    return default


def normalize_word(text: str) -> str:
    lowered = text.replace("’", "'").lower()
    trimmed = re.sub(r"^[^a-z']+|[^a-z']+$", "", lowered)
    trimmed = CONTRACTION_DIRECT_MAP.get(trimmed, trimmed)

    if trimmed.endswith("'re") and len(trimmed) > 3:
        trimmed = trimmed[:-3]
    elif trimmed.endswith("'m") and len(trimmed) > 2:
        trimmed = trimmed[:-2]
    elif trimmed.endswith("'ve") and len(trimmed) > 3:
        trimmed = trimmed[:-3]
    elif trimmed.endswith("'ll") and len(trimmed) > 3:
        trimmed = trimmed[:-3]

    return trimmed


def simple_suffix_variants(word: str) -> list[str]:
    variants = []

    if len(word) > 4 and word.endswith("ies"):
        variants.append(word[:-3] + "y")

    if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
        variants.append(word[:-1])

    if len(word) > 4 and word.endswith("es"):
        variants.append(word[:-2])

    if len(word) > 5 and word.endswith("ing"):
        base = word[:-3]
        variants.append(base)
        variants.append(base + "e")

    if len(word) > 4 and word.endswith("ed"):
        base = word[:-2]
        variants.append(base)
        variants.append(base + "e")

    return [item for item in variants if item and re.fullmatch(r"[a-z']+", item)]


def lemma_candidates(text: str) -> list[str]:
    normalized = normalize_word(text)
    if not normalized:
        return []

    candidates = [normalized]

    simplemma_lemma = simplemma.lemmatize(normalized, lang="en")
    if simplemma_lemma:
        candidates.append(simplemma_lemma.lower())

    for item in list(candidates):
        candidates.extend(simple_suffix_variants(item))

    unique = []
    seen = set()
    for item in candidates:
        item = item.lower()
        if not item or item in seen:
            continue
        seen.add(item)
        unique.append(item)

    return unique


def primary_lemma(text: str) -> str:
    candidates = lemma_candidates(text)
    return candidates[0] if candidates else ""


def normalize_entry(entry: dict, order: int, default_unit: str = "Imported") -> dict | None:
    raw = entry.get("lemma") or entry.get("word")
    lemma = primary_lemma(str(raw or ""))
    if not lemma:
        return None

    locations = entry.get("locations") if isinstance(entry.get("locations"), list) else []
    if not locations:
        locations = [{"unit": safe_text(entry.get("unit")) or default_unit, "page": safe_int(entry.get("page"))}]

    return {
        "lemma": lemma,
        "unit": safe_text(locations[0].get("unit")) or default_unit,
        "page": safe_int(locations[0].get("page")),
        "locations": locations,
        "frequency": normalize_frequency(entry.get("frequency")),
        "in_syllabus": bool(entry.get("in_syllabus", True)),
        "meaning": safe_text(entry.get("meaning")),
        "order": int(entry.get("order", order)),
        "customFields": entry.get("customFields") if isinstance(entry.get("customFields"), dict) else {},
    }


def dedupe_entries_keep_first(entries: list[dict]) -> list[dict]:
    seen = set()
    result = []
    for index, entry in enumerate(entries):
        normalized = normalize_entry(entry, index)
        if not normalized:
            continue
        location = normalized["locations"][0]
        key = (normalized["lemma"], location.get("unit"), location.get("page"))
        if key in seen:
            continue
        seen.add(key)
        result.append(normalized)
    return result


def migrate_old_vocab_files() -> None:
    libraries = read_json(LIBRARIES_PATH, [])
    if not isinstance(libraries, list):
        libraries = []

    changed = False

    has_migrated_book = any(lib.get("id") == "legacy_book_vocab" for lib in libraries)
    if OLD_BOOK_VOCAB_PATH.exists() and not has_migrated_book:
        old_entries = read_json(OLD_BOOK_VOCAB_PATH, [])
        if isinstance(old_entries, list) and old_entries:
            entries = dedupe_entries_keep_first(old_entries)
            libraries.append(
                {
                    "id": "legacy_book_vocab",
                    "name": "旧主词库 book_vocab",
                    "type": "main",
                    "createdAt": now_iso(),
                    "entries": entries,
                }
            )
            changed = True

    has_extra = any(lib.get("id") == EXTRA_LIBRARY_ID for lib in libraries)
    if OLD_EXTRA_VOCAB_PATH.exists() and not has_extra:
        old_entries = read_json(OLD_EXTRA_VOCAB_PATH, [])
        if isinstance(old_entries, list):
            entries = dedupe_entries_keep_first(old_entries)
            libraries.append(
                {
                    "id": EXTRA_LIBRARY_ID,
                    "name": "补充词库",
                    "type": "extra",
                    "createdAt": now_iso(),
                    "entries": entries,
                }
            )
            changed = True

    if not any(lib.get("id") == EXTRA_LIBRARY_ID for lib in libraries):
        libraries.append(
            {
                "id": EXTRA_LIBRARY_ID,
                "name": "补充词库",
                "type": "extra",
                "createdAt": now_iso(),
                "entries": [],
            }
        )
        changed = True

    if not any(lib.get("id") == BASIC_LIBRARY_ID for lib in libraries):
        basic_entries = []
        seen = set()
        for word in sorted(BUILTIN_STOPWORDS):
            lemma = primary_lemma(word)
            if not lemma or lemma in seen:
                continue
            seen.add(lemma)
            basic_entries.append(
                {
                    "lemma": lemma,
                    "unit": "Basic",
                    "page": None,
                    "frequency": "low",
                    "in_syllabus": False,
                    "meaning": "基础功能词",
                    "order": len(basic_entries),
                }
            )

        libraries.append(
            {
                "id": BASIC_LIBRARY_ID,
                "name": "基础词白名单",
                "type": "basic",
                "createdAt": now_iso(),
                "entries": basic_entries,
            }
        )
        changed = True

    if changed:
        write_json(LIBRARIES_PATH, libraries)


def get_extra_library(libraries: list[dict]) -> dict:
    for lib in libraries:
        if lib.get("id") == EXTRA_LIBRARY_ID:
            return lib

    extra = {
        "id": EXTRA_LIBRARY_ID,
        "name": "补充词库",
        "type": "extra",
        "createdAt": now_iso(),
        "entries": [],
    }
    libraries.append(extra)
    return extra


def get_extra_libraries(libraries: list[dict]) -> list[dict]:
    extras = [library for library in libraries if library.get("type") == "extra"]
    if extras:
        return extras
    return [get_extra_library(libraries)]


def get_default_extra_library_id(libraries: list[dict]) -> str:
    extras = get_extra_libraries(libraries)
    configured = safe_text(sqlite_load_setting("default_extra_library_id", EXTRA_LIBRARY_ID))
    if any(library.get("id") == configured for library in extras):
        return configured
    fallback = extras[0].get("id", EXTRA_LIBRARY_ID)
    sqlite_save_setting("default_extra_library_id", fallback)
    return fallback


def get_extra_library_by_id(libraries: list[dict], library_id: str | None = None) -> dict:
    target_id = safe_text(library_id) or get_default_extra_library_id(libraries)
    target = next(
        (library for library in libraries if library.get("id") == target_id and library.get("type") == "extra"),
        None,
    )
    if not target:
        raise HTTPException(status_code=404, detail="补充词库不存在。")
    return target


def get_basic_library(libraries: list[dict]) -> dict:
    for lib in libraries:
        if lib.get("id") == BASIC_LIBRARY_ID:
            return lib

    basic = {
        "id": BASIC_LIBRARY_ID,
        "name": "基础词白名单",
        "type": "basic",
        "createdAt": now_iso(),
        "entries": [],
    }
    libraries.append(basic)
    return basic


def current_basic_stopwords() -> set[str]:
    libraries = load_libraries()
    basic_library = get_basic_library(libraries)
    words = set(BUILTIN_STOPWORDS)

    for entry in basic_library.get("entries", []):
        lemma = str(entry.get("lemma", "")).strip().lower()
        if lemma:
            words.add(lemma)

    return words


def make_library_id() -> str:
    return f"library_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:8]}"


def decode_text_file(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=400, detail="文本文件编码无法识别，请改存为 UTF-8 后重试。")


def extract_entries_from_text(
    text: str,
    default_unit: str = "Imported",
    default_frequency: str = "low",
    default_in_syllabus: bool = True,
) -> list[dict]:
    matches = re.findall(r"[A-Za-z]+(?:'[A-Za-z]+)?", text)
    entries = []
    seen = set()

    for raw_word in matches:
        lemma = primary_lemma(raw_word)
        if not lemma or lemma in seen:
            continue
        seen.add(lemma)
        entries.append(
            {
                "lemma": lemma,
                "unit": default_unit,
                "page": None,
                "frequency": default_frequency,
                "in_syllabus": default_in_syllabus,
                "meaning": None,
                "order": len(entries),
            }
        )

    return entries


def pick_column_index(headers: list[str], aliases: set[str]) -> int | None:
    for index, header in enumerate(headers):
        if header in aliases:
            return index
    return None


def get_cell(row: list, index: int | None):
    if index is None or index >= len(row):
        return None
    return row[index]


def parse_table_rows(headers: list, rows: list[list], default_unit: str = "Imported") -> list[dict]:
    normalized_headers = [normalize_header(header) for header in headers]

    word_index = pick_column_index(normalized_headers, WORD_COLUMN_ALIASES)
    unit_index = pick_column_index(normalized_headers, UNIT_COLUMN_ALIASES)
    page_index = pick_column_index(normalized_headers, PAGE_COLUMN_ALIASES)
    frequency_index = pick_column_index(normalized_headers, FREQUENCY_COLUMN_ALIASES)
    in_syllabus_index = pick_column_index(normalized_headers, IN_SYLLABUS_COLUMN_ALIASES)
    oversyllabus_index = pick_column_index(normalized_headers, OVERSYLLABUS_COLUMN_ALIASES)
    meaning_index = pick_column_index(normalized_headers, MEANING_COLUMN_ALIASES)

    if word_index is None:
        return []

    entries = []
    for row in rows:
        raw_word = safe_text(get_cell(row, word_index))
        if not raw_word:
            continue

        lemma = primary_lemma(raw_word)
        if not lemma:
            continue

        if in_syllabus_index is not None:
            in_syllabus = parse_bool(get_cell(row, in_syllabus_index), True)
        elif oversyllabus_index is not None:
            in_syllabus = not parse_bool(get_cell(row, oversyllabus_index), False)
        else:
            in_syllabus = True

        custom_fields = {}
        standard_indices = {word_index, unit_index, page_index, frequency_index, in_syllabus_index, oversyllabus_index, meaning_index}
        for index, header in enumerate(headers):
            value = safe_text(get_cell(row, index))
            if index not in standard_indices and value:
                custom_fields[str(header).strip() or f"column_{index + 1}"] = value

        entries.append(
            {
                "lemma": lemma,
                "unit": safe_text(get_cell(row, unit_index)) or default_unit,
                "page": safe_int(get_cell(row, page_index)),
                "frequency": normalize_frequency(get_cell(row, frequency_index)),
                "in_syllabus": in_syllabus,
                "meaning": safe_text(get_cell(row, meaning_index)),
                "order": len(entries),
                "customFields": custom_fields,
            }
        )

    return entries


def parse_csv_text(text: str) -> list[dict]:
    sample = "\n".join(text.splitlines()[:10])
    delimiter = ","

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [row for row in reader if any(str(cell).strip() for cell in row)]

    if not rows:
        return []

    headers = rows[0]
    data_rows = rows[1:]
    return parse_table_rows(headers, data_rows)


def parse_xlsx_content(content: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    worksheet = workbook.active
    raw_rows = list(worksheet.iter_rows(values_only=True))

    rows = []
    for raw_row in raw_rows:
        row = list(raw_row)
        if any(safe_text(cell) for cell in row):
            rows.append(row)

    if not rows:
        return []

    headers = rows[0]
    data_rows = rows[1:]
    return parse_table_rows(headers, data_rows)


def parse_upload_table(filename: str, content: bytes) -> tuple[list[str], list[list]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".txt"}:
        text = decode_text_file(content)
        sample = "\n".join(text.splitlines()[:10])
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = [row for row in reader if any(safe_text(cell) for cell in row)]
        return (list(map(str, rows[0])) if rows else [], rows[1:] if rows else [])
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        rows = [list(row) for row in workbook.active.iter_rows(values_only=True) if any(safe_text(cell) for cell in row)]
        return (list(map(str, rows[0])) if rows else [], rows[1:] if rows else [])
    entries = parse_uploaded_vocab(filename, content)
    keys = ["lemma", "unit", "page", "frequency", "in_syllabus", "meaning"]
    return (keys, [[item.get(key) for key in keys] for item in entries])


def infer_import_mapping(headers: list[str]) -> dict[str, str]:
    aliases = {
        "lemma": WORD_COLUMN_ALIASES,
        "meaning": MEANING_COLUMN_ALIASES,
        "unit": UNIT_COLUMN_ALIASES,
        "lesson": LESSON_COLUMN_ALIASES,
        "page": PAGE_COLUMN_ALIASES,
        "frequency": FREQUENCY_COLUMN_ALIASES,
        "pos": POS_COLUMN_ALIASES,
        "ipa": IPA_COLUMN_ALIASES,
        "audio_url": AUDIO_URL_COLUMN_ALIASES,
        "serial": SERIAL_COLUMN_ALIASES,
        "in_syllabus": IN_SYLLABUS_COLUMN_ALIASES,
    }
    mapping = {}
    for field, values in aliases.items():
        for header in headers:
            if normalize_header(header) in values:
                mapping[field] = header
                break
    return mapping


def entries_from_mapping(headers: list[str], rows: list[list], mapping: dict[str, str], field_definitions: list[dict] | None = None) -> list[dict]:
    header_map = {str(header): index for index, header in enumerate(headers)}
    mapped = {field: header_map.get(str(source)) for field, source in mapping.items()}
    word_index = mapped.get("lemma")
    if word_index is None:
        return []
    entries = []
    for row in rows:
        raw_word = safe_text(get_cell(row, word_index))
        lemma = primary_lemma(raw_word or "")
        if not lemma:
            continue
        known_sources = {index for index in mapped.values() if index is not None}
        for definition in field_definitions or []:
            source = safe_text(definition.get("sourceHeader") or definition.get("source")) if isinstance(definition, dict) else ""
            if source in header_map:
                known_sources.add(header_map[source])
        fields = {str(header): safe_text(get_cell(row, index)) for index, header in enumerate(headers) if index not in known_sources and safe_text(get_cell(row, index))}
        for field in ("lesson", "pos", "ipa", "audio_url", "serial"):
            value = safe_text(get_cell(row, mapped.get(field)))
            if value:
                fields[dict(IMPORT_MAPPING_FIELDS)[field]] = value
        for definition in field_definitions or []:
            if not isinstance(definition, dict):
                continue
            key = safe_text(definition.get("key"))
            source = safe_text(definition.get("sourceHeader") or definition.get("source"))
            if not key or not source or source not in header_map:
                continue
            value = safe_text(get_cell(row, header_map[source]))
            if value:
                fields[key] = value
        unit = safe_text(get_cell(row, mapped.get("unit")))
        lesson = safe_text(get_cell(row, mapped.get("lesson")))
        entries.append({
            "lemma": lemma,
            "unit": unit or lesson or "Imported",
            "page": safe_int(get_cell(row, mapped.get("page"))),
            "frequency": normalize_frequency(get_cell(row, mapped.get("frequency"))),
            "in_syllabus": parse_bool(get_cell(row, mapped.get("in_syllabus")), True),
            "meaning": safe_text(get_cell(row, mapped.get("meaning"))),
            "order": len(entries),
            "customFields": fields,
        })
    return entries


def parse_uploaded_vocab(filename: str, content: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise HTTPException(status_code=400, detail="仅支持 csv、txt、xlsx、docx、pdf。")

    if suffix == ".csv":
        text = decode_text_file(content)
        return parse_csv_text(text)

    if suffix == ".xlsx":
        try:
            return parse_xlsx_content(content)
        except Exception as error:
            raise HTTPException(status_code=400, detail=f"XLSX 解析失败：{error}")

    if suffix == ".txt":
        text = decode_text_file(content)
        first_line = next((line for line in text.splitlines() if line.strip()), "")
        if "," in first_line or "\t" in first_line or ";" in first_line:
            try:
                return parse_csv_text(text)
            except Exception:
                pass
        return extract_entries_from_text(text)

    if suffix == ".docx":
        try:
            doc = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            return extract_entries_from_text(text)
        except Exception as error:
            raise HTTPException(status_code=400, detail=f"DOCX 解析失败：{error}")

    if suffix == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return extract_entries_from_text(text)
        except Exception as error:
            raise HTTPException(status_code=400, detail=f"PDF 解析失败：{error}")

    return []


def parse_chat_upload_file(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix not in AI_CHAT_SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="AI 对话暂支持 txt、md、csv、json、docx、pdf、xlsx、py、js、ts、jsx、tsx、html、css、xml、yaml、yml、log。",
        )

    if suffix in AI_CHAT_TEXT_EXTENSIONS:
        return decode_text_file(content)

    if suffix == ".docx":
        doc = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        worksheet = workbook.active
        rows = []
        for raw_row in worksheet.iter_rows(values_only=True):
            row = [str(cell).strip() for cell in raw_row if cell is not None and str(cell).strip()]
            if row:
                rows.append("\t".join(row))
        return "\n".join(rows)

    return ""


def summarize_library(lib: dict) -> dict:
    return {
        "id": lib.get("id"),
        "name": lib.get("name"),
        "type": lib.get("type"),
        "createdAt": lib.get("createdAt"),
        "count": len(lib.get("entries", [])),
        "displayConfig": normalized_display_config(lib),
        "availableFields": available_display_fields(lib),
    }


def merge_library_entries(source_libraries: list[dict]) -> list[dict]:
    merged: dict[str, dict] = {}
    for library in source_libraries:
        for source in sorted(library.get("entries", []), key=lambda item: int(item.get("order", 0))):
            entry = normalize_entry(source, len(merged))
            if not entry:
                continue
            lemma = entry["lemma"]
            if lemma not in merged:
                merged[lemma] = copy.deepcopy(entry)
                continue
            target = merged[lemma]
            known_locations = {(item.get("unit"), item.get("page")) for item in target.get("locations", [])}
            for location in entry.get("locations", []):
                location_key = (location.get("unit"), location.get("page"))
                if location_key not in known_locations:
                    target.setdefault("locations", []).append(location)
                    known_locations.add(location_key)
            target["unit"] = target["locations"][0].get("unit") or "Imported"
            target["page"] = target["locations"][0].get("page")
            if not target.get("meaning") and entry.get("meaning"):
                target["meaning"] = entry["meaning"]
            if entry.get("frequency") == "high":
                target["frequency"] = "high"
            target["in_syllabus"] = bool(target.get("in_syllabus")) or bool(entry.get("in_syllabus"))
            for key, value in entry.get("customFields", {}).items():
                if key and value is not None and not target.setdefault("customFields", {}).get(key):
                    target["customFields"][key] = value
    result = list(merged.values())
    for order, entry in enumerate(result):
        entry["order"] = order
    return result


def get_selected_libraries(
    selected_library_ids: list[str],
    include_extra: bool,
) -> list[dict]:
    libraries = load_libraries()
    selected = []
    id_set = set(selected_library_ids or [])

    for lib in libraries:
        lib_id = lib.get("id")
        lib_type = lib.get("type")
        if lib_type == "main" and lib_id in id_set:
            selected.append(lib)
        if include_extra and lib_type == "extra":
            selected.append(lib)

    return selected


def build_vocab_index_from_libraries(libraries: list[dict]) -> dict[str, dict]:
    index: dict[str, dict] = {}

    for library_order, lib in enumerate(libraries):
        lib_id = lib.get("id")
        lib_name = lib.get("name")
        lib_type = lib.get("type")

        entries = lib.get("entries", [])
        for entry in sorted(entries, key=lambda item: int(item.get("order", 0))):
            lemma = str(entry.get("lemma", "")).lower().strip()
            if not lemma or lemma in index:
                continue

            index[lemma] = {
                **entry,
                "lemma": lemma,
                "libraryId": lib_id,
                "libraryName": lib_name,
                "libraryType": lib_type,
                "libraryOrder": library_order,
            }

    return index


def get_merged_words(
    library_ids: list[str] | None = None,
    include_extra: bool = False,
    include_basic: bool = False,
) -> list[dict]:
    libraries = load_libraries()
    result = []
    seen = set()
    id_set = set(library_ids or [])

    for library_order, lib in enumerate(libraries):
        lib_id = lib.get("id")
        lib_type = lib.get("type")

        should_include = False
        if lib_type == "main":
            if library_ids is None:
                should_include = True
            else:
                should_include = lib_id in id_set

        if include_extra and lib_type == "extra":
            should_include = True

        if include_basic and lib_id == BASIC_LIBRARY_ID:
            should_include = True

        if not should_include:
            continue

        entries = sorted(lib.get("entries", []), key=lambda item: int(item.get("order", 0)))

        for entry in entries:
            lemma = str(entry.get("lemma", "")).strip().lower()
            if not lemma or lemma in seen:
                continue
            seen.add(lemma)
            result.append(
                {
                    **entry,
                    "lemma": lemma,
                    "libraryId": lib_id,
                    "libraryName": lib.get("name"),
                    "libraryType": lib_type,
                    "libraryOrder": library_order,
                    "frequencyText": frequency_text(entry.get("frequency", "low")),
                }
            )

    return result


def get_level(status: dict, lemma: str) -> int:
    try:
        level = int(status.get(lemma, {}).get("level", 0))
    except Exception:
        level = 0
    return min(max(level, 0), 12)


def level_group(level: int) -> str:
    if level <= 0:
        return "untracked"
    if level == 1:
        return "unknown"
    if level == 2:
        return "fuzzy"
    if 3 <= level <= 5:
        return "known"
    if 6 <= level <= 11:
        return "stableish"
    return "stable"


def level_label(level: int) -> str:
    if level <= 0:
        return "未追踪"
    if level == 1:
        return "不认识"
    if level == 2:
        return "模糊"
    if 3 <= level <= 5:
        return "认识"
    if 6 <= level <= 11:
        return "较稳定"
    return "稳定"


def match_level_filter(level: int, level_filter: str) -> bool:
    if level_filter == "all":
        return True
    if level_filter in {"0", "1", "2", "12"}:
        return level == int(level_filter)
    if level_filter == "known":
        return 3 <= level <= 5
    if level_filter == "stableish":
        return 6 <= level <= 11
    return True


def filter_words(
    words: list[dict],
    query: str,
    level_filter: str,
    status: dict,
) -> list[dict]:
    q = (query or "").strip().lower()
    result = []

    for word in words:
        lemma = word.get("lemma", "")
        level = get_level(status, lemma)

        if not match_level_filter(level, level_filter):
            continue

        if q:
            locations_text = " ".join(
                f"{location.get('unit', '')} {location.get('page', '')}"
                for location in (word.get("locations") or [])
                if isinstance(location, dict)
            )
            custom_fields_text = " ".join(
                f"{key} {value}" for key, value in (word.get("customFields") or {}).items()
            )
            searchable = " ".join(
                [
                    str(word.get("lemma", "")),
                    str(word.get("unit", "")),
                    str(word.get("page", "")),
                    str(word.get("frequency", "")),
                    str(word.get("libraryName", "")),
                    str(word.get("libraryType", "")),
                    str(word.get("meaning", "")),
                    locations_text,
                    custom_fields_text,
                ]
            ).lower()
            if q not in searchable:
                continue

        result.append(
            {
                **word,
                "level": level,
                "levelGroup": level_group(level),
                "levelLabel": level_label(level),
                "frequencyText": frequency_text(word.get("frequency", "low")),
            }
        )

    return result


def paginate(items: list[dict], page: int, page_size: int) -> dict:
    total = len(items)
    page_size = min(max(page_size, 1), 100)
    total_pages = max((total + page_size - 1) // page_size, 1)
    page = min(max(page, 1), total_pages)
    start = (page - 1) * page_size
    end = start + page_size

    return {
        "items": items[start:end],
        "page": page,
        "pageSize": page_size,
        "total": total,
        "totalPages": total_pages,
    }


def build_mastery_overview(words: list[dict] | None = None) -> dict:
    status = load_user_status()
    levels = {str(i): 0 for i in range(13)}

    if words is None:
        words = get_merged_words(library_ids=None, include_extra=True, include_basic=True)

    seen = set()
    for word in words:
        lemma = word.get("lemma")
        if not lemma or lemma in seen:
            continue
        seen.add(lemma)
        level = get_level(status, lemma)
        levels[str(level)] += 1

    return {
        "levels": levels,
        "trackedWords": len(seen),
    }


def outside_text(analysis_mode: AnalysisMode) -> str:
    return "超纲" if analysis_mode == "strict" else "词库外"


def choose_worst_mark(existing: MarkState | None, new_mark: MarkState) -> MarkState:
    priority = {
        "unknown": 0,
        "fuzzy": 1,
        "known": 2,
    }
    if existing is None:
        return new_mark
    return existing if priority[existing] <= priority[new_mark] else new_mark


def choose_focus_mark(fuzzy_count: int, unknown_count: int) -> str:
    if fuzzy_count > 0:
        return "fuzzy"
    if unknown_count > 0:
        return "unknown"
    return "known"


def update_level_by_mark(
    previous_level: int,
    previous_mark: str | None,
    current_mark: MarkState,
) -> int:
    level = previous_level

    if current_mark == "known":
        if level < 3:
            level = 3
        if previous_mark == "known":
            level += 1
    elif current_mark == "fuzzy":
        if previous_mark == "known":
            level -= 1
        elif level < 2:
            level = 2
    else:
        if previous_mark == "known":
            level -= 2
        elif level < 1:
            level = 1

    return min(max(level, 0), 12)


def find_vocab_entry(vocab_index: dict[str, dict], token_text: str) -> tuple[str, dict | None]:
    candidates = lemma_candidates(token_text)
    for candidate in candidates:
        if candidate in vocab_index:
            return candidate, vocab_index[candidate]

    if candidates:
        return candidates[0], None

    return "", None


def focus_priority(item: dict) -> tuple:
    library_type = item.get("libraryType")
    focus_mark = item.get("focusMark")
    freq = item.get("frequency")
    count = int(item.get("count", 0))
    matched = bool(item.get("matchedInVocab"))

    if library_type == "main" and focus_mark == "fuzzy":
        return (0, -count, item.get("lemma", ""))
    if library_type == "main" and focus_mark == "unknown":
        return (1, -count, item.get("lemma", ""))
    if library_type == "main" and freq == "high":
        return (2, -count, item.get("lemma", ""))
    if library_type == "extra" and focus_mark in {"fuzzy", "unknown"}:
        return (3, -count, item.get("lemma", ""))
    if not matched and count >= 2:
        return (4, -count, item.get("lemma", ""))
    if not matched:
        return (5, -count, item.get("lemma", ""))
    return (9, -count, item.get("lemma", ""))


def mask_api_key(api_key: str) -> str:
    if not api_key:
        return ""
    if len(api_key) <= 8:
        return "*" * len(api_key)
    return f"{api_key[:4]}{'*' * max(len(api_key) - 8, 4)}{api_key[-4:]}"


def merge_ai_settings(override: AISettingsPayload | None) -> dict:
    saved = load_ai_settings()
    merged = saved.copy()
    payload = override.model_dump(exclude_none=True) if override is not None else {}
    merged.update({key: value for key, value in payload.items() if key != "apiKey"})

    # Request keys are session-only. Persisted settings deliberately never contain API keys.
    merged["apiKey"] = safe_text(payload.get("apiKey")) or environment_api_key()

    merged["provider"] = safe_text(merged.get("provider")) or "openai-compatible"
    merged["baseUrl"] = safe_text(merged.get("baseUrl")) or ""
    merged["model"] = safe_text(merged.get("model")) or ""
    merged["apiKey"] = safe_text(merged.get("apiKey")) or ""

    try:
        merged["temperature"] = float(merged.get("temperature", 0.2))
    except Exception:
        merged["temperature"] = 0.2

    try:
        merged["timeoutSeconds"] = int(merged.get("timeoutSeconds", 30))
    except Exception:
        merged["timeoutSeconds"] = 30

    merged["timeoutSeconds"] = min(max(merged["timeoutSeconds"], 5), 120)
    merged["temperature"] = min(max(merged["temperature"], 0.0), 2.0)
    merged["verifySSL"] = bool(merged.get("verifySSL", True))

    return merged


def parse_settings_json(settings_json: str) -> dict:
    settings_json = safe_text(settings_json)
    if not settings_json:
        return load_ai_settings()

    try:
        data = json.loads(settings_json)
    except Exception:
        raise HTTPException(status_code=400, detail="AI settingsJson 不是有效 JSON。")

    if not isinstance(data, dict):
        raise HTTPException(status_code=400, detail="AI settingsJson 结构不正确。")

    payload = AISettingsPayload(**{
        "provider": data.get("provider", "openai-compatible"),
        "baseUrl": data.get("baseUrl", ""),
        "model": data.get("model", ""),
        "apiKey": data.get("apiKey", ""),
        "temperature": data.get("temperature", 0.2),
        "timeoutSeconds": data.get("timeoutSeconds", 30),
        "verifySSL": data.get("verifySSL", True),
    })

    return merge_ai_settings(payload)


def validate_ai_settings(settings: dict) -> None:
    if not settings.get("baseUrl"):
        raise HTTPException(status_code=400, detail="AI Base URL 未配置。")
    if not settings.get("model"):
        raise HTTPException(status_code=400, detail="AI Model 未配置。")
    if not settings.get("apiKey"):
        raise HTTPException(
            status_code=400,
            detail="AI API Key 未配置。请设置 AI_API_KEY 环境变量，或在当前页面临时输入 Key。",
        )


def ai_chat_endpoint(base_url: str) -> str:
    base_url = base_url.strip()
    if base_url.endswith("/chat/completions"):
        return base_url
    return f"{base_url.rstrip('/')}/chat/completions"


def sha_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def word_cache_key(settings: dict, lemma: str) -> str:
    raw = f"{settings.get('provider')}|{settings.get('baseUrl')}|{settings.get('model')}|word|{lemma}"
    return sha_text(raw)


def article_cache_key(settings: dict, text: str) -> str:
    raw = f"{settings.get('provider')}|{settings.get('baseUrl')}|{settings.get('model')}|article|{text}"
    return sha_text(raw)


def chat_cache_key(settings: dict, message: str, file_texts: list[dict], history: list[dict] | None = None) -> str:
    joined_files = "\n".join(
        f"[{item['name']}]\n{item['content']}"
        for item in file_texts
    )
    history_text = json.dumps((history or [])[-10:], ensure_ascii=False, sort_keys=True)
    raw = f"{settings.get('provider')}|{settings.get('baseUrl')}|{settings.get('model')}|chat|{message}|{joined_files}|{history_text}"
    return sha_text(raw)


def extract_json_string(text: str) -> str | None:
    json_block_patterns = [
        r"```json\s*(\{.*?\})\s*```",
        r"```\s*(\{.*?\})\s*```",
    ]

    for pattern in json_block_patterns:
        match = re.search(pattern, text, flags=re.DOTALL)
        if match:
            return match.group(1)

    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        return text[start:end + 1]

    return None


def parse_word_ai_json(raw_text: str) -> dict:
    try:
        direct = json.loads(raw_text)
        if isinstance(direct, dict):
            return direct
    except Exception:
        pass

    json_text = extract_json_string(raw_text)
    if not json_text:
        raise HTTPException(status_code=502, detail="AI 返回内容无法解析为 JSON。")

    try:
        data = json.loads(json_text)
    except Exception:
        raise HTTPException(status_code=502, detail="AI 返回 JSON 解析失败。")

    if not isinstance(data, dict):
        raise HTTPException(status_code=502, detail="AI 返回 JSON 结构不正确。")

    return data


def normalize_pos_detail_item(item) -> dict | None:
    if isinstance(item, dict):
        english = safe_text(item.get("en") or item.get("english") or item.get("pos"))
        chinese = safe_text(item.get("zh") or item.get("chinese") or item.get("meaning"))
    else:
        raw = safe_text(item)
        if not raw:
            return None
        lower = raw.lower()
        english = None
        chinese = None

        for pos_en, pos_zh in POS_ZH_MAP.items():
            if pos_en in lower or raw == pos_zh:
                english = pos_en
                chinese = pos_zh
                break

        if not english:
            english = raw
        if not chinese:
            chinese = POS_ZH_MAP.get(english.lower(), "")

    if not english and not chinese:
        return None

    return {
        "en": english or "",
        "zh": chinese or "",
    }


def normalize_pos_details(value) -> list[dict]:
    if isinstance(value, str):
        value = [item.strip() for item in re.split(r"[,，/;；]+", value) if item.strip()]
    if not isinstance(value, list):
        value = []

    result = []
    seen = set()

    for item in value:
        normalized = normalize_pos_detail_item(item)
        if not normalized:
            continue

        key = f"{normalized['en']}|{normalized['zh']}"
        if key in seen:
            continue

        seen.add(key)
        result.append(normalized)

    return result


def normalize_derivative_item(item) -> dict | None:
    if isinstance(item, dict):
        word = safe_text(item.get("word") or item.get("lemma") or item.get("name"))
        meaning = safe_text(item.get("meaning") or item.get("translation"))
        relation = safe_text(item.get("relation") or item.get("form") or item.get("type"))
        source = safe_text(item.get("source")) or "ai"
    else:
        raw = safe_text(item)
        if not raw:
            return None

        parts = re.split(r"[:：]\s*", raw, maxsplit=1)
        word = safe_text(parts[0])
        meaning = safe_text(parts[1]) if len(parts) > 1 else None
        relation = None
        source = "ai"

    if not word and not meaning:
        return None

    return {
        "word": word or "",
        "meaning": meaning or "",
        "relation": relation or "",
        "source": source,
    }


def normalize_derivative_items(value) -> list[dict]:
    if isinstance(value, str):
        value = [item.strip() for item in re.split(r"[,，/;；]+", value) if item.strip()]
    if not isinstance(value, list):
        value = []

    result = []
    seen = set()

    for item in value:
        normalized = normalize_derivative_item(item)
        if not normalized:
            continue

        key = f"{normalized['word']}|{normalized['meaning']}|{normalized['relation']}"
        if key in seen:
            continue

        seen.add(key)
        result.append(normalized)

    return result


def ai_frequency_text(value: str) -> str:
    return "高考高频" if value == "high" else "高考低频"


def normalize_ai_word_result(word: str, data: dict) -> dict:
    pos_details = normalize_pos_details(data.get("posDetails", data.get("pos", [])))
    derivatives = normalize_derivative_items(data.get("derivatives", []))

    ai_frequency = str(data.get("aiFrequency", "")).strip().lower()
    if "high" in ai_frequency or "高" in ai_frequency:
        ai_frequency = "high"
    else:
        ai_frequency = "low"

    normalized_word = safe_text(data.get("word")) or word
    lemma = (
        safe_text(data.get("lemma"))
        or primary_lemma(normalized_word)
        or primary_lemma(word)
        or normalized_word
    )

    return {
        "word": normalized_word,
        "lemma": lemma,
        "meaning": safe_text(data.get("meaning")) or "",
        "posDetails": pos_details,
        "pos": [item["en"] or item["zh"] for item in pos_details if item["en"] or item["zh"]],
        "ipa": safe_text(data.get("ipa") or data.get("phonetic")) or "",
        "audioUrl": safe_text(data.get("audioUrl") or data.get("audio_url")) or "",
        "derivatives": derivatives,
        "aiFrequency": ai_frequency,
        "aiFrequencyText": ai_frequency_text(ai_frequency),
        "example": safe_text(data.get("example")) or "",
        "exampleTranslation": safe_text(data.get("exampleTranslation")) or "",
    }


def build_ssl_context(verify_ssl: bool):
    if verify_ssl:
        return ssl.create_default_context()
    return ssl._create_unverified_context()


def call_ai_chat(settings: dict, messages: list[dict]) -> str:
    validate_ai_settings(settings)

    endpoint = ai_chat_endpoint(settings["baseUrl"])

    payload = {
        "model": settings["model"],
        "messages": messages,
        "temperature": settings.get("temperature", 0.2),
    }

    request_body = json.dumps(payload).encode("utf-8")

    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {settings['apiKey']}",
    }

    req = urllib.request.Request(
        endpoint,
        data=request_body,
        headers=headers,
        method="POST",
    )

    context = build_ssl_context(bool(settings.get("verifySSL", True)))

    try:
        with urllib.request.urlopen(
            req,
            timeout=settings.get("timeoutSeconds", 30),
            context=context,
        ) as response:
            body = response.read().decode("utf-8")
    except urllib.error.HTTPError as error:
        try:
            error_body = error.read().decode("utf-8")
        except Exception:
            error_body = str(error)
        raise HTTPException(status_code=502, detail=f"AI 接口错误：{error_body}")
    except urllib.error.URLError as error:
        detail = str(error)
        if "CERTIFICATE_VERIFY_FAILED" in detail:
            raise HTTPException(
                status_code=502,
                detail="AI 请求失败：SSL 证书校验失败。可在 AI 设置中关闭 verifySSL 后重试，或改用有效证书接口。",
            )
        raise HTTPException(status_code=502, detail=f"AI 请求失败：{error}")
    except TimeoutError:
        raise HTTPException(status_code=504, detail="AI 请求超时。")

    try:
        data = json.loads(body)
    except Exception:
        raise HTTPException(status_code=502, detail="AI 返回内容不是有效 JSON。")

    content = (
        data.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
    )

    if not content:
        raise HTTPException(status_code=502, detail="AI 未返回有效内容。")

    return content


def build_word_query_messages(word: str, context: str | None, library_context: dict | None = None) -> list[dict]:
    context_block = f"文章上下文：{context}" if context else "无上下文"
    library_context = library_context if isinstance(library_context, dict) else {}
    library_bits = []
    for label, key in (("词库", "libraryName"), ("教材位置", "unit"), ("页码", "page"), ("已有释义", "meaning"), ("自定义信息", "customFields")):
        value = library_context.get(key)
        if value not in (None, "", {}, []):
            library_bits.append(f"{label}：{json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value}")
    library_block = "\n".join(library_bits) if library_bits else "无词库补充信息"
    return [
        {
            "role": "system",
            "content": (
                "你是英语词汇学习助手。"
                "请严格只返回 JSON，不要返回 Markdown，不要返回解释文字。"
                "JSON 结构必须为："
                '{"word":"","lemma":"","meaning":"","ipa":"","audioUrl":"","posDetails":[{"en":"","zh":""}],"derivatives":[{"word":"","meaning":"","relation":"","source":"ai"}],"aiFrequency":"high|low","example":"","exampleTranslation":""}'
            ),
        },
        {
            "role": "user",
            "content": (
                f"请分析这个英语单词：{word}\n"
                f"{context_block}\n"
                f"词库上下文：\n{library_block}\n"
                "请给出：原词、lemma、中文翻译、国际音标（没有把握则空字符串）、词性（英文+中文）、常见派生词及每个派生词的中文翻译。派生词必须带 relation，例如‘过去式 +ed’或‘名词派生’，source 固定为 ai。audioUrl 仅在有可靠的直接音频 URL 时填写，否则为空字符串。不要猜测词根。"
            ),
        },
    ]


def build_article_translate_messages(text: str) -> list[dict]:
    return [
        {
            "role": "system",
            "content": (
                "你是英语文章翻译助手。"
                "请把用户给出的英文自然翻译成中文，保留原文段落结构。"
                "不要添加额外分析，不要编号，不要解释。"
            ),
        },
        {
            "role": "user",
            "content": text,
        },
    ]


def build_import_draft_messages(text: str) -> list[dict]:
    return [
        {
            "role": "system",
            "content": (
                "你是词库导入预处理助手。"
                "请从半结构化词表中提取英语词条。"
                "严格只返回 JSON，不要 Markdown 或解释。"
                "JSON 结构必须为："
                '{"headers":["word","meaning"],"rows":[["example","示例"]]}。'
                "headers 可使用 word、meaning、unit、lesson、page、frequency、pos、ipa、audio_url、serial；"
                "仅保留可确认的英文单词。"
            ),
        },
        {"role": "user", "content": text[:20000]},
    ]


def build_chat_messages(message: str, file_texts: list[dict], history: list[dict] | None = None) -> list[dict]:
    file_block = ""
    if file_texts:
        chunks = []
        for item in file_texts:
            chunks.append(f"文件名：{item['name']}\n文件内容：\n{item['content']}")
        file_block = "\n\n".join(chunks)

    user_content = message.strip()
    if file_block:
        user_content += f"\n\n以下是用户上传的文件内容，请结合一起回答：\n{file_block}"

    messages = [
        {
            "role": "system",
            "content": (
                "你是英语学习和阅读辅助 AI。"
                "请根据用户问题与上传文件内容回答。"
                "回答要准确、简洁、结构清晰。"
            ),
        },
    ]
    for item in (history or [])[-10:]:
        role = item.get("role")
        content = safe_text(item.get("content"))
        if role in {"user", "assistant"} and content:
            messages.append({"role": role, "content": content})
    messages.append({"role": "user", "content": user_content})
    return messages


@app.on_event("startup")
def startup() -> None:
    ensure_data_files()


@app.get("/api/health")
def health():
    ensure_data_files()
    libraries = load_libraries()
    return {
        "ok": True,
        "libraries": len(libraries),
    }


@app.get("/api/libraries")
def list_libraries():
    ensure_data_files()
    libraries = load_libraries()
    return {
        "libraries": [summarize_library(lib) for lib in libraries],
        "defaultExtraLibraryId": get_default_extra_library_id(libraries),
    }


@app.get("/api/libraries/{library_id}/display-config")
def get_library_display_config(library_id: str):
    libraries = load_libraries()
    library = next((item for item in libraries if item.get("id") == library_id), None)
    if not library:
        raise HTTPException(status_code=404, detail="词库不存在。")
    return {"libraryId": library_id, "displayConfig": normalized_display_config(library)}


@app.put("/api/libraries/{library_id}/display-config")
def update_library_display_config(library_id: str, request: DisplayConfigRequest):
    libraries = load_libraries()
    library = next((item for item in libraries if item.get("id") == library_id), None)
    if not library:
        raise HTTPException(status_code=404, detail="词库不存在。")
    library["displayConfig"] = validate_display_config(library, request.displayConfig)
    save_libraries(libraries)
    return {"ok": True, "libraryId": library_id, "displayConfig": library["displayConfig"]}


@app.get("/api/basic-words")
def list_basic_words():
    ensure_data_files()
    libraries = load_libraries()
    basic_library = get_basic_library(libraries)
    entries = sorted(basic_library.get("entries", []), key=lambda item: int(item.get("order", 0)))
    return {
        "library": summarize_library(basic_library),
        "words": entries,
    }


@app.post("/api/basic-words/add")
def add_basic_words(request: AddBasicWordsRequest):
    ensure_data_files()
    if not request.words:
        raise HTTPException(status_code=400, detail="请至少添加一个单词。")

    libraries = load_libraries()
    basic_library = get_basic_library(libraries)
    entries = basic_library.get("entries", [])

    existing = {str(entry.get("lemma", "")).strip().lower() for entry in entries}
    added = []

    for raw_word in request.words:
        lemma = primary_lemma(raw_word)
        if not lemma or lemma in existing:
            continue
        existing.add(lemma)
        entries.append(
            {
                "lemma": lemma,
                "unit": "Basic",
                "page": None,
                "frequency": "low",
                "in_syllabus": False,
                "meaning": "基础功能词",
                "order": len(entries),
            }
        )
        added.append(lemma)

    basic_library["entries"] = entries
    save_libraries(libraries)

    return {
        "ok": True,
        "added": added,
        "count": len(added),
        "library": summarize_library(basic_library),
    }


@app.post("/api/basic-words/delete")
def delete_basic_words(request: BasicWordsDeleteRequest):
    ensure_data_files()
    targets = {primary_lemma(word) for word in request.words if primary_lemma(word)}
    if not targets:
        raise HTTPException(status_code=400, detail="请至少指定一个有效单词。")
    libraries = load_libraries()
    basic_library = get_basic_library(libraries)
    previous = basic_library.get("entries", [])
    basic_library["entries"] = [entry for entry in previous if str(entry.get("lemma", "")).lower() not in targets]
    for index, entry in enumerate(basic_library["entries"]):
        entry["order"] = index
    save_libraries(libraries)
    return {"ok": True, "deleted": len(previous) - len(basic_library["entries"]), "library": summarize_library(basic_library)}


@app.post("/api/basic-words/reset")
def reset_basic_words():
    ensure_data_files()
    libraries = load_libraries()
    basic_library = get_basic_library(libraries)
    basic_library["entries"] = []
    save_libraries(libraries)
    return {"ok": True, "library": summarize_library(basic_library)}


@app.get("/api/basic-words/export")
def export_basic_words():
    ensure_data_files()
    libraries = load_libraries()
    basic_library = get_basic_library(libraries)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=["lemma", "meaning"])
    writer.writeheader()
    for entry in basic_library.get("entries", []):
        writer.writerow({"lemma": entry.get("lemma"), "meaning": entry.get("meaning") or "基础功能词"})
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=basic_whitelist.csv"})


@app.get("/api/ai/providers")
def get_ai_providers():
    ensure_data_files()
    return {
        "providers": AI_PROVIDER_PRESETS,
    }


@app.get("/api/ai/settings")
def get_ai_settings():
    ensure_data_files()
    settings = merge_ai_settings(None)
    return {
        "provider": settings.get("provider", "openai-compatible"),
        "baseUrl": settings.get("baseUrl", ""),
        "model": settings.get("model", ""),
        "apiKeyMasked": mask_api_key(settings.get("apiKey", "")),
        "hasApiKey": bool(settings.get("apiKey")),
        "temperature": settings.get("temperature", 0.2),
        "timeoutSeconds": settings.get("timeoutSeconds", 30),
        "verifySSL": bool(settings.get("verifySSL", True)),
        "providers": AI_PROVIDER_PRESETS,
        "fieldHelp": {
            "temperature": "控制回答随机性。越低越稳定，推荐 0.1~0.3；越高越发散。",
            "timeoutSeconds": "单次 AI 请求最长等待秒数。超过后会报超时。",
            "verifySSL": "是否校验 HTTPS 证书。证书报错时可临时关闭，但仅建议本地测试使用。",
        },
    }


@app.post("/api/ai/settings")
def save_ai_settings_endpoint(payload: AISettingsPayload):
    ensure_data_files()
    merged = merge_ai_settings(payload)
    save_ai_settings(merged)
    return {
        "ok": True,
        "provider": merged.get("provider"),
        "baseUrl": merged.get("baseUrl"),
        "model": merged.get("model"),
        "apiKeyMasked": mask_api_key(merged.get("apiKey", "")),
        "hasApiKey": bool(merged.get("apiKey")),
        "temperature": merged.get("temperature"),
        "timeoutSeconds": merged.get("timeoutSeconds"),
        "verifySSL": merged.get("verifySSL"),
    }


@app.post("/api/ai/word-query")
def ai_word_query(request: AIWordQueryRequest):
    ensure_data_files()

    lemma = primary_lemma(request.word)
    if not lemma:
        raise HTTPException(status_code=400, detail="无效单词。")

    settings = merge_ai_settings(request.settings)
    cache = load_ai_cache()
    context_signature = json.dumps(request.libraryContext or {}, ensure_ascii=False, sort_keys=True) + (request.context or "")
    cache_key = word_cache_key(settings, lemma + "|" + hashlib.sha256(context_signature.encode("utf-8")).hexdigest()[:12])

    if not request.forceRefresh and cache_key in cache["word"]:
        cached = cache["word"][cache_key]
        normalized = normalize_ai_word_result(lemma, cached.get("data", {}))

        if normalized != cached.get("data"):
            cache["word"][cache_key]["data"] = normalized
            save_ai_cache(cache)

        return {
            "ok": True,
            "cached": True,
            "data": normalized,
            "updatedAt": cached.get("updatedAt"),
        }

    messages = build_word_query_messages(lemma, request.context, request.libraryContext)
    raw_content = call_ai_chat(settings, messages)
    parsed = parse_word_ai_json(raw_content)
    normalized = normalize_ai_word_result(lemma, parsed)

    cache["word"][cache_key] = {
        "type": "word",
        "word": lemma,
        "provider": settings.get("provider"),
        "baseUrl": settings.get("baseUrl"),
        "model": settings.get("model"),
        "updatedAt": now_iso(),
        "data": normalized,
    }
    save_ai_cache(cache)

    return {
        "ok": True,
        "cached": False,
        "data": normalized,
        "updatedAt": cache["word"][cache_key]["updatedAt"],
    }


@app.post("/api/vocab/ai-import-draft")
def ai_import_draft(request: AIImportDraftRequest):
    ensure_data_files()
    text = request.text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="请先提供要整理的词表文本。")
    if len(text) > 20000:
        raise HTTPException(status_code=400, detail="待整理文本不能超过 20000 个字符。")

    raw_content = call_ai_chat(merge_ai_settings(request.settings), build_import_draft_messages(text))
    parsed = parse_word_ai_json(raw_content)

    headers = [str(item).strip() for item in parsed.get("headers", []) if str(item).strip()]
    raw_rows = parsed.get("rows", [])
    if not headers or not isinstance(raw_rows, list):
        raise HTTPException(status_code=502, detail="AI 草稿缺少 headers 或 rows。")

    rows = []
    for raw_row in raw_rows:
        if isinstance(raw_row, list):
            rows.append(([safe_text(cell) or "" for cell in raw_row[:len(headers)]] + [""] * len(headers))[:len(headers)])

    mapping = infer_import_mapping(headers)
    entries = entries_from_mapping(headers, rows, mapping)
    return {
        "headers": headers,
        "rows": rows,
        "mapping": mapping,
        "mappingFields": [{"key": key, "label": label} for key, label in IMPORT_MAPPING_FIELDS],
        "availableMappingFields": [{"key": key, "label": label} for key, label in IMPORT_MAPPING_FIELDS],
        "rowCount": len(rows),
        "recognizedCount": len(entries),
        "message": "AI 已生成草稿，请检查映射和内容后再确认导入。",
    }


@app.post("/api/vocab/ai-import-file-draft")
async def ai_import_file_draft(
    file: UploadFile = File(...),
    settingsJson: str = Form("{}"),
):
    ensure_data_files()
    filename = file.filename or ""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空。")
    if len(content) > MAX_IMPORT_FILE_BYTES:
        raise HTTPException(status_code=413, detail="导入文件不能超过 5 MB。")
    try:
        settings = AISettingsPayload.model_validate(json.loads(settingsJson or "{}"))
    except Exception:
        raise HTTPException(status_code=400, detail="AI 设置格式不正确。")
    try:
        text = parse_chat_upload_file(filename, content)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"文件内容提取失败：{error}")
    if not text.strip():
        raise HTTPException(status_code=400, detail="未能从文件中提取可供 AI 整理的文本。")
    draft = ai_import_draft(AIImportDraftRequest(text=text[:20000], settings=settings))
    draft["filename"] = filename
    draft["message"] = "AI 已根据上传文件生成草稿；请确认字段映射和内容后再导入。"
    return draft


@app.post("/api/ai/article-translate")
def ai_article_translate(request: AIArticleTranslateRequest):
    ensure_data_files()

    text = safe_text(request.text)
    if not text:
        raise HTTPException(status_code=400, detail="文章内容为空。")

    settings = merge_ai_settings(request.settings)
    cache = load_ai_cache()
    cache_key = article_cache_key(settings, text)

    if not request.forceRefresh and cache_key in cache["article"]:
        cached = cache["article"][cache_key]
        return {
            "ok": True,
            "cached": True,
            "translation": cached["translation"],
            "updatedAt": cached.get("updatedAt"),
        }

    messages = build_article_translate_messages(text)
    translation = call_ai_chat(settings, messages).strip()

    cache["article"][cache_key] = {
        "type": "article",
        "provider": settings.get("provider"),
        "baseUrl": settings.get("baseUrl"),
        "model": settings.get("model"),
        "updatedAt": now_iso(),
        "translation": translation,
    }
    save_ai_cache(cache)

    return {
        "ok": True,
        "cached": False,
        "translation": translation,
        "updatedAt": cache["article"][cache_key]["updatedAt"],
    }


@app.get("/api/chat/sessions")
def get_chat_sessions():
    ensure_data_files()
    return {"sessions": list_chat_sessions()}


@app.post("/api/chat/sessions")
def create_new_chat_session(request: ChatSessionCreateRequest):
    ensure_data_files()
    title = safe_text(request.title) or "新对话"
    return {"session": create_chat_session(f"chat-{uuid.uuid4().hex}", title)}


@app.get("/api/chat/sessions/{session_id}")
def get_one_chat_session(session_id: str):
    ensure_data_files()
    session = get_chat_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="对话不存在。")
    return {"session": session}


@app.put("/api/chat/sessions/{session_id}")
def rename_one_chat_session(session_id: str, request: ChatSessionRenameRequest):
    ensure_data_files()
    if not get_chat_session(session_id):
        raise HTTPException(status_code=404, detail="对话不存在。")
    title = safe_text(request.title) or "未命名对话"
    rename_chat_session(session_id, title)
    return {"ok": True, "title": title}


@app.delete("/api/chat/sessions/{session_id}")
def delete_one_chat_session(session_id: str):
    ensure_data_files()
    if not delete_chat_session(session_id):
        raise HTTPException(status_code=404, detail="对话不存在。")
    return {"ok": True}


@app.post("/api/ai/chat")
async def ai_chat(
    message: str = Form(...),
    forceRefresh: str = Form("false"),
    settingsJson: str = Form(""),
    sessionId: str = Form(""),
    files: list[UploadFile] | None = File(None),
):
    ensure_data_files()

    clean_message = safe_text(message)
    if not clean_message:
        raise HTTPException(status_code=400, detail="聊天内容为空。")

    settings = parse_settings_json(settingsJson)
    clean_session_id = safe_text(sessionId)
    session = get_chat_session(clean_session_id) if clean_session_id else None
    if clean_session_id and not session:
        raise HTTPException(status_code=404, detail="对话不存在，请刷新后重试。")
    file_texts = []

    if files:
        for file in files:
            filename = file.filename or "unnamed"
            content = await file.read()
            if not content:
                continue
            extracted = parse_chat_upload_file(filename, content)
            file_texts.append({
                "name": filename,
                "content": extracted[:20000],
            })

    history = session.get("messages", []) if session else []
    cache = load_ai_cache()
    force_refresh_flag = coerce_bool(forceRefresh, False)
    cache_key = chat_cache_key(settings, clean_message, file_texts, history)

    if not force_refresh_flag and cache_key in cache["chat"]:
        cached = cache["chat"][cache_key]
        if clean_session_id:
            attached_files = [{"name": item["name"]} for item in file_texts]
            append_chat_message(clean_session_id, "user", clean_message, attached_files)
            append_chat_message(clean_session_id, "assistant", cached["reply"], cached.get("files", []))
        return {
            "ok": True,
            "cached": True,
            "reply": cached["reply"],
            "updatedAt": cached.get("updatedAt"),
            "files": cached.get("files", []),
        }

    messages = build_chat_messages(clean_message, file_texts, history)
    reply = call_ai_chat(settings, messages).strip()

    cache["chat"][cache_key] = {
        "type": "chat",
        "provider": settings.get("provider"),
        "baseUrl": settings.get("baseUrl"),
        "model": settings.get("model"),
        "updatedAt": now_iso(),
        "reply": reply,
        "files": [{"name": item["name"]} for item in file_texts],
    }
    save_ai_cache(cache)

    if clean_session_id:
        append_chat_message(clean_session_id, "user", clean_message, [{"name": item["name"]} for item in file_texts])
        append_chat_message(clean_session_id, "assistant", reply, cache["chat"][cache_key]["files"])

    return {
        "ok": True,
        "cached": False,
        "reply": reply,
        "updatedAt": cache["chat"][cache_key]["updatedAt"],
        "files": cache["chat"][cache_key]["files"],
    }


@app.get("/api/mastery/overview")
def mastery_overview(
    libraryIds: str = Query(""),
    includeExtra: bool = Query(True),
):
    ensure_data_files()

    ids = [item for item in libraryIds.split(",") if item]
    if ids:
        words = get_merged_words(library_ids=ids, include_extra=includeExtra, include_basic=False)
    else:
        words = get_merged_words(library_ids=None, include_extra=includeExtra, include_basic=False)

    return build_mastery_overview(words)


@app.get("/api/libraries/merged/words")
def get_merged_library_words(
    libraryIds: str = Query(""),
    includeExtra: bool = Query(False),
    includeBasic: bool = Query(False),
    page: int = Query(1),
    pageSize: int = Query(50),
    query: str = Query(""),
    levelFilter: str = Query("all"),
):
    ensure_data_files()

    ids = [item for item in libraryIds.split(",") if item]
    if ids:
        words = get_merged_words(library_ids=ids, include_extra=includeExtra, include_basic=includeBasic)
    else:
        words = get_merged_words(library_ids=None, include_extra=includeExtra, include_basic=includeBasic)

    status = load_user_status()
    filtered = filter_words(words, query, levelFilter, status)
    page_data = paginate(filtered, page, pageSize)

    return {
        **page_data,
        "overview": build_mastery_overview(words),
    }


@app.get("/api/libraries/{library_id}/words")
def get_library_words(
    library_id: str,
    page: int = Query(1),
    pageSize: int = Query(50),
    query: str = Query(""),
    levelFilter: str = Query("all"),
):
    ensure_data_files()

    libraries = load_libraries()
    target = next((lib for lib in libraries if lib.get("id") == library_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="词库不存在。")

    words = []
    for entry in sorted(target.get("entries", []), key=lambda item: int(item.get("order", 0))):
        words.append(
            {
                **entry,
                "libraryId": target.get("id"),
                "libraryName": target.get("name"),
                "libraryType": target.get("type"),
                "frequencyText": frequency_text(entry.get("frequency", "low")),
            }
        )

    status = load_user_status()
    filtered = filter_words(words, query, levelFilter, status)
    page_data = paginate(filtered, page, pageSize)

    return {
        **page_data,
        "overview": build_mastery_overview(words),
    }


@app.post("/api/vocab/upload")
async def upload_vocab(
    file: UploadFile = File(...),
    libraryName: str = Form(""),
):
    ensure_data_files()

    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise HTTPException(status_code=400, detail="仅支持 csv、txt、xlsx、docx、pdf。")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空。")

    entries = parse_uploaded_vocab(filename, content)
    if not entries:
        raise HTTPException(
            status_code=400,
            detail="没有识别到可导入的英文单词。推荐优先使用 csv / xlsx。",
        )

    libraries = load_libraries()
    name = libraryName.strip() or Path(filename).stem or "未命名词库"

    new_library = {
        "id": make_library_id(),
        "name": name,
        "type": "main",
        "createdAt": now_iso(),
        "entries": entries,
    }

    libraries.append(new_library)
    get_extra_library(libraries)
    get_basic_library(libraries)
    save_libraries(libraries)

    return {
        "ok": True,
        "library": summarize_library(new_library),
        "importedCount": len(entries),
        "message": "词库上传成功。",
        "libraries": [summarize_library(lib) for lib in libraries],
    }


@app.post("/api/vocab/import-preview")
async def vocab_import_preview(file: UploadFile = File(...)):
    ensure_data_files()
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空。")
    if len(content) > MAX_IMPORT_FILE_BYTES:
        raise HTTPException(status_code=413, detail="导入文件不能超过 5 MB。")
    headers, rows = parse_upload_table(file.filename or "", content)
    mapping = infer_import_mapping(headers)
    entries = entries_from_mapping(headers, rows, mapping)
    return {
        "filename": file.filename or "", "headers": headers, "rows": rows,
        "mapping": mapping, "sample": entries[:10], "rowCount": len(rows),
        "recognizedCount": len(entries),
        "mappingFields": [{"key": key, "label": label} for key, label in IMPORT_MAPPING_FIELDS],
        "fieldDefinitions": [],
        "message": "请确认单词列映射后再导入；单元、页码、频率等均为可选，未映射列会作为自定义字段保留。",
    }


@app.post("/api/vocab/import-confirm")
def vocab_import_confirm(request: ImportConfirmRequest):
    ensure_data_files()
    field_definitions = []
    used_keys = set()
    for definition in request.fieldDefinitions:
        if not isinstance(definition, dict):
            continue
        label = safe_text(definition.get("label"))[:80]
        key = re.sub(r"[^A-Za-z0-9_-]+", "_", safe_text(definition.get("key"))).strip("_")[:80]
        source_header = safe_text(definition.get("sourceHeader") or definition.get("source"))
        if not label or not key or source_header not in request.headers or key in used_keys:
            continue
        used_keys.add(key)
        field_definitions.append({
            "key": key,
            "label": label,
            "sourceHeader": source_header,
            "displaySource": f"customFields.{key}",
        })
    entries = entries_from_mapping(request.headers, request.rows, request.mapping, field_definitions)
    if not entries:
        raise HTTPException(status_code=400, detail="未识别到单词；请将一个源列映射为“单词/lemma”。")
    libraries = load_libraries()
    if request.targetType == "basic":
        basic_library = get_basic_library(libraries)
        existing = {str(entry.get("lemma", "")).lower() for entry in basic_library.get("entries", [])}
        added = 0
        for entry in entries:
            if entry["lemma"] in existing:
                continue
            existing.add(entry["lemma"])
            basic_library.setdefault("entries", []).append({
                **entry,
                "unit": "Basic",
                "locations": [{"unit": "Basic", "page": None}],
                "in_syllabus": False,
                "meaning": entry.get("meaning") or "基础功能词",
                "order": len(basic_library["entries"]),
            })
            added += 1
        save_libraries(libraries)
        return {"ok": True, "library": summarize_library(basic_library), "importedCount": added}
    library = {
        "id": make_library_id(), "name": safe_text(request.libraryName) or "未命名词库",
        "type": "main", "createdAt": now_iso(), "entries": entries,
        "fieldDefinitions": field_definitions,
    }
    libraries.append(library)
    get_extra_library(libraries)
    get_basic_library(libraries)
    save_libraries(libraries)
    return {"ok": True, "library": summarize_library(library), "importedCount": len(entries)}


@app.get("/api/libraries/{library_id}/export")
def export_library(library_id: str):
    ensure_data_files()
    target = next((lib for lib in load_libraries() if lib.get("id") == library_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="词库不存在。")
    output = io.StringIO()
    fieldnames = ["lemma", "unit", "page", "frequency", "in_syllabus", "meaning", "locations", "custom_fields"]
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    for entry in target.get("entries", []):
        writer.writerow({
            "lemma": entry.get("lemma"), "unit": entry.get("unit"), "page": entry.get("page"),
            "frequency": entry.get("frequency"), "in_syllabus": entry.get("in_syllabus"),
            "meaning": entry.get("meaning") or "", "locations": json.dumps(entry.get("locations", []), ensure_ascii=False),
            "custom_fields": json.dumps(entry.get("customFields", {}), ensure_ascii=False),
        })
    filename = f"{target.get('name', 'vocab')}.csv"
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(filename)}"})


@app.delete("/api/libraries/{library_id}")
def delete_library(library_id: str, confirm: str = Query("")):
    ensure_data_files()
    if confirm != "DELETE":
        raise HTTPException(status_code=400, detail="请确认删除操作。")
    libraries = load_libraries()
    target = next((lib for lib in libraries if lib.get("id") == library_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="词库不存在。")
    if target.get("type") == "basic":
        raise HTTPException(status_code=400, detail="基础词白名单不能删除。")
    if target.get("type") == "extra":
        extras = get_extra_libraries(libraries)
        if len(extras) <= 1:
            raise HTTPException(status_code=400, detail="至少需要保留一个补充词库。")
        if library_id == get_default_extra_library_id(libraries):
            raise HTTPException(status_code=400, detail="请先在词库管理中切换默认补充词库，再删除当前词库。")
    save_libraries([lib for lib in libraries if lib.get("id") != library_id])
    return {"ok": True, "deleted": summarize_library(target)}


@app.post("/api/extra-libraries")
def create_extra_library(request: LibraryCreateRequest):
    ensure_data_files()
    name = safe_text(request.name)
    if not name:
        raise HTTPException(status_code=400, detail="请填写补充词库名称。")
    libraries = load_libraries()
    library = {"id": make_library_id(), "name": name, "type": "extra", "createdAt": now_iso(), "entries": []}
    libraries.append(library)
    save_libraries(libraries)
    return {"ok": True, "library": summarize_library(library)}


@app.put("/api/libraries/{library_id}")
def rename_library(library_id: str, request: LibraryRenameRequest):
    ensure_data_files()
    name = safe_text(request.name)
    if not name:
        raise HTTPException(status_code=400, detail="请填写词库名称。")
    libraries = load_libraries()
    target = next((library for library in libraries if library.get("id") == library_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="词库不存在。")
    if target.get("type") == "basic":
        raise HTTPException(status_code=400, detail="基础词白名单不能重命名。")
    target["name"] = name
    save_libraries(libraries)
    return {"ok": True, "library": summarize_library(target)}


@app.post("/api/libraries/merge")
def merge_libraries(request: LibraryMergeRequest):
    ensure_data_files()
    source_ids = list(dict.fromkeys(item for item in request.sourceLibraryIds if safe_text(item)))
    name = safe_text(request.name)
    if len(source_ids) != 2:
        raise HTTPException(status_code=400, detail="请选择两个不同的主词库。")
    if not name:
        raise HTTPException(status_code=400, detail="请填写合并后词库名称。")
    libraries = load_libraries()
    sources = [next((library for library in libraries if library.get("id") == source_id), None) for source_id in source_ids]
    if any(source is None or source.get("type") != "main" for source in sources):
        raise HTTPException(status_code=400, detail="只能合并两个主词库。")
    merged_library = {"id": make_library_id(), "name": name, "type": "main", "createdAt": now_iso(), "entries": merge_library_entries(sources)}
    next_libraries = [library for library in libraries if not request.deleteSources or library.get("id") not in source_ids]
    next_libraries.append(merged_library)
    save_libraries(next_libraries)
    return {"ok": True, "library": summarize_library(merged_library), "sourceDeleted": bool(request.deleteSources)}


@app.post("/api/extra-libraries/default")
def set_default_extra_library(request: DefaultExtraLibraryRequest):
    ensure_data_files()
    libraries = load_libraries()
    target = get_extra_library_by_id(libraries, request.libraryId)
    sqlite_save_setting("default_extra_library_id", target["id"])
    return {"ok": True, "defaultExtraLibraryId": target["id"]}


@app.get("/api/review/due")
def get_review_due(limit: int = Query(30, ge=1, le=100)):
    ensure_data_files()
    status = load_user_status()
    words = get_merged_words(library_ids=None, include_extra=True, include_basic=False)
    candidates = filter_words(words, "", "all", status)
    due_schedules = get_due_review_schedules(now_iso())
    candidates = [word for word in candidates if 0 < word.get("level", 0) < 12]
    for word in candidates:
        schedule = due_schedules.get(word.get("lemma")) or get_review_schedule(word.get("lemma"))
        word["dueAt"] = schedule.get("dueAt")
        word["intervalDays"] = schedule.get("intervalDays", 0)
    candidates = [word for word in candidates if word.get("dueAt") is None or word.get("lemma") in due_schedules]
    candidates.sort(key=lambda word: (word.get("dueAt") is not None, word.get("level", 0), word.get("lemma", "")))
    return {"items": candidates[:limit], "total": len(candidates)}


@app.post("/api/review/mark")
def mark_review(request: ReviewMarkRequest):
    ensure_data_files()
    lemma = primary_lemma(request.lemma)
    if not lemma:
        raise HTTPException(status_code=400, detail="无效单词。")
    status = load_user_status()
    previous = status.get(lemma, {})
    level = update_level_by_mark(int(previous.get("level", 0)), previous.get("lastMark"), request.mark)
    status[lemma] = {**previous, "level": level, "lastMark": request.mark, "updatedAt": now_iso(), "seenCount": int(previous.get("seenCount", 0)) + 1, "knownStreak": int(previous.get("knownStreak", 0)) + 1 if request.mark == "known" else 0}
    save_user_status(status)
    schedule = get_review_schedule(lemma)
    old_interval = int(schedule.get("intervalDays", 0))
    old_repetitions = int(schedule.get("repetitions", 0))
    if request.mark == "known":
        interval = min(max(old_interval * 2, 1), 90)
        repetitions = old_repetitions + 1
    elif request.mark == "fuzzy":
        interval = 1
        repetitions = max(old_repetitions - 1, 0)
    else:
        interval = 0
        repetitions = 0
    due_at = (datetime.now(timezone.utc) + timedelta(days=interval)).isoformat()
    save_review_schedule(lemma, due_at, interval, repetitions)
    return {"ok": True, "lemma": lemma, "level": level, "levelLabel": level_label(level), "dueAt": due_at, "intervalDays": interval}


@app.post("/api/user-status/reset")
def reset_user_status():
    ensure_data_files()
    save_user_status({})
    clear_review_schedules()
    return {
        "ok": True,
        "message": "认识情况已重置。",
        "overview": build_mastery_overview(),
    }


# The Vite development server owns the frontend in source mode. Packaged desktop
# builds bundle frontend/dist as frontend_dist and are served by this FastAPI app.
if FRONTEND_DIST_DIR.is_dir():
    frontend_assets_dir = FRONTEND_DIST_DIR / "assets"
    if frontend_assets_dir.is_dir():
        app.mount("/assets", StaticFiles(directory=frontend_assets_dir), name="frontend-assets")

    @app.get("/", include_in_schema=False)
    def desktop_frontend_index():
        return FileResponse(FRONTEND_DIST_DIR / "index.html")

    @app.get("/{frontend_path:path}", include_in_schema=False)
    def desktop_frontend_fallback(frontend_path: str):
        if frontend_path.startswith("api/"):
            raise HTTPException(status_code=404, detail="接口不存在。")
        return FileResponse(FRONTEND_DIST_DIR / "index.html")


@app.post("/api/extra-vocab/add")
def add_extra_vocab(request: AddExtraWordRequest):
    ensure_data_files()

    lemma = primary_lemma(request.word)
    if not lemma:
        raise HTTPException(status_code=400, detail="无效单词。")

    libraries = load_libraries()
    extra = get_extra_library_by_id(libraries, request.libraryId)
    entries = extra.get("entries", [])

    exists = any(str(entry.get("lemma", "")).lower() == lemma for entry in entries)

    if not exists:
        entries.append(
            {
                "lemma": lemma,
                "unit": "Supplement",
                "page": None,
                "frequency": request.frequency,
                "in_syllabus": False,
                "meaning": safe_text(request.meaning),
                "order": len(entries),
            }
        )

    extra["entries"] = entries
    save_libraries(libraries)

    status = load_user_status()
    current = status.get(lemma, {})
    level = min(max(int(request.level), 1), 12)

    status[lemma] = {
        **current,
        "level": max(int(current.get("level", 0)), level),
        "lastMark": current.get("lastMark"),
        "updatedAt": now_iso(),
        "seenCount": int(current.get("seenCount", 0)),
        "knownStreak": int(current.get("knownStreak", 0)),
    }
    save_user_status(status)

    return {
        "ok": True,
        "lemma": lemma,
        "library": summarize_library(extra),
        "overview": build_mastery_overview(),
    }


@app.post("/api/analyze")
def analyze(request: AnalyzeRequest):
    ensure_data_files()

    if not request.selectedLibraryIds and not request.includeExtra:
        raise HTTPException(status_code=400, detail="请先选择至少一个词库。")

    selected_libraries = get_selected_libraries(
        selected_library_ids=request.selectedLibraryIds,
        include_extra=request.includeExtra,
    )

    if not selected_libraries:
        raise HTTPException(status_code=400, detail="没有可用词库，请先上传词库或选择补充词库。")

    vocab_index = build_vocab_index_from_libraries(selected_libraries)

    if not vocab_index:
        raise HTTPException(status_code=400, detail="所选词库为空，请先上传有效词库。")

    stopword_set = current_basic_stopwords()

    word_tokens = [token for token in request.tokens if token.isWord]

    if not word_tokens:
        raise HTTPException(status_code=400, detail="没有检测到可分析的英文单词。")

    if len(word_tokens) > MAX_ANALYZE_WORDS:
        raise HTTPException(status_code=400, detail="文章太长，请缩短到 5000 个英文词以内。")

    status = load_user_status()

    unique_lemmas = set()
    lemma_final_marks: dict[str, MarkState] = {}
    analyzed_tokens = []
    focus_map: dict[str, dict] = {}
    gloss_map: dict[str, str] = {}

    stats = {
        "totalTokens": 0,
        "uniqueWords": 0,
        "recognized": 0,
        "fuzzy": 0,
        "unmastered": 0,
        "unknown": 0,
        "mainHits": 0,
        "extraHits": 0,
        "outsideVocab": 0,
        "ignoredBasics": 0,
        "mainFocus": 0,
        "bookFocus": 0,
        "oversyllabus": 0,
    }

    for token in request.tokens:
        if not token.isWord:
            analyzed_tokens.append(
                {
                    "id": token.id,
                    "text": token.text,
                    "isWord": False,
                }
            )
            continue

        stats["totalTokens"] += 1

        mark: MarkState = request.tokenMarks.get(str(token.id), "unknown")
        lemma, entry = find_vocab_entry(vocab_index, token.text)
        normalized = primary_lemma(token.text)
        final_lemma = lemma or normalized or ""

        if final_lemma:
            unique_lemmas.add(final_lemma)

        matched = entry is not None
        library_type = entry.get("libraryType") if entry else "untracked"
        frequency = entry.get("frequency", "low") if entry else "low"
        meaning = entry.get("meaning") if entry else None
        is_basic_ignored = bool(request.ignoreBasicWords and final_lemma in stopword_set)
        outside = not matched
        outside_label = outside_text(request.analysisMode)

        if matched and final_lemma:
            previous = lemma_final_marks.get(final_lemma)
            lemma_final_marks[final_lemma] = choose_worst_mark(previous, mark)

        if meaning and final_lemma and final_lemma not in gloss_map:
            gloss_map[final_lemma] = meaning

        current_level = get_level(status, final_lemma) if final_lemma else 0

        if is_basic_ignored:
            stats["ignoredBasics"] += 1
        else:
            if mark == "known":
                stats["recognized"] += 1
            elif mark == "fuzzy":
                stats["fuzzy"] += 1
            else:
                stats["unmastered"] += 1
                stats["unknown"] += 1

            if matched:
                if library_type == "main":
                    stats["mainHits"] += 1
                elif library_type == "extra":
                    stats["extraHits"] += 1
            else:
                stats["outsideVocab"] += 1
                stats["oversyllabus"] += 1

            if matched and library_type == "main" and mark != "known":
                stats["mainFocus"] += 1
                stats["bookFocus"] += 1

        if not is_basic_ignored and mark != "known":
            focus_key = final_lemma or token.text.lower()
            if focus_key:
                if focus_key not in focus_map:
                    focus_map[focus_key] = {
                        "lemma": focus_key,
                        "displayText": token.text,
                        "count": 0,
                        "fuzzyCount": 0,
                        "unknownCount": 0,
                        "matchedInVocab": matched,
                        "libraryName": entry.get("libraryName") if entry else None,
                        "libraryType": library_type,
                        "unit": entry.get("unit") if entry else None,
                        "page": entry.get("page") if entry else None,
                        "frequency": frequency,
                        "frequencyText": frequency_text(frequency),
                        "meaning": meaning,
                        "locations": entry.get("locations", []) if entry else [],
                        "customFields": entry.get("customFields", {}) if entry else {},
                        "outsideText": outside_label,
                        "outside": outside,
                        "basicIgnored": False,
                        "canAddToExtra": not matched,
                        "suggestedLevel": 2 if mark == "fuzzy" else 1,
                    }

                focus_map[focus_key]["count"] += 1
                if mark == "fuzzy":
                    focus_map[focus_key]["fuzzyCount"] += 1
                if mark == "unknown":
                    focus_map[focus_key]["unknownCount"] += 1

        analyzed_tokens.append(
            {
                "id": token.id,
                "text": token.text,
                "isWord": True,
                "selectionMark": mark,
                "normalized": normalized,
                "lemma": final_lemma,
                "level": current_level,
                "levelLabel": level_label(current_level),
                "matchedInVocab": matched,
                "libraryName": entry.get("libraryName") if entry else None,
                "libraryType": library_type,
                "sourceText": "补充词库" if library_type == "extra" else (entry.get("libraryName") if entry else None),
                "unit": entry.get("unit") if entry else None,
                "page": entry.get("page") if entry else None,
                "frequency": frequency,
                "frequencyText": frequency_text(frequency) if matched else None,
                "meaning": meaning,
                "locations": entry.get("locations", []) if entry else [],
                "customFields": entry.get("customFields", {}) if entry else {},
                "outside": outside,
                "outsideText": outside_label,
                "basicIgnored": is_basic_ignored,
                "basicText": "基础词" if is_basic_ignored else None,
                "canAddToExtra": not matched,
            }
        )

    stats["uniqueWords"] = len(unique_lemmas)

    for lemma_key, final_mark in lemma_final_marks.items():
        current = status.get(lemma_key, {})
        previous_level = get_level(status, lemma_key)
        previous_mark = current.get("lastMark")

        next_level = update_level_by_mark(previous_level, previous_mark, final_mark)

        seen_count = int(current.get("seenCount", 0)) + 1
        known_streak = int(current.get("knownStreak", 0))

        if final_mark == "known":
            known_streak += 1
        else:
            known_streak = 0

        status[lemma_key] = {
            "level": next_level,
            "lastMark": final_mark,
            "updatedAt": now_iso(),
            "seenCount": seen_count,
            "knownStreak": known_streak,
        }

    save_user_status(status)

    updated_analyzed_tokens = []
    for item in analyzed_tokens:
        if item.get("isWord") and item.get("lemma"):
            new_level = get_level(status, item["lemma"])
            item["level"] = new_level
            item["levelLabel"] = level_label(new_level)
        updated_analyzed_tokens.append(item)

    focus_words = []
    for item in focus_map.values():
        focus_mark = choose_focus_mark(item["fuzzyCount"], item["unknownCount"])
        focus_words.append(
            {
                **item,
                "focusMark": focus_mark,
            }
        )

    focus_words = sorted(focus_words, key=focus_priority)

    selected_words = get_merged_words(
        library_ids=request.selectedLibraryIds,
        include_extra=request.includeExtra,
        include_basic=False,
    )

    article_gloss_lines = []
    for lemma_key, meaning in gloss_map.items():
        if meaning:
            article_gloss_lines.append(f"{lemma_key}：{meaning}")

    return {
        "tokens": updated_analyzed_tokens,
        "stats": stats,
        "focusWords": focus_words,
        "masteryOverview": build_mastery_overview(selected_words),
        "analysisMode": request.analysisMode,
        "ignoreBasicWords": request.ignoreBasicWords,
        "articleGlossaryText": "\n".join(article_gloss_lines[:80]),
    }